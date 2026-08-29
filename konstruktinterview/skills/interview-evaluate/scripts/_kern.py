"""Gemeinsamer Kern der Auswertungsskripte.

Dieses Modul wird von ``build_workbook.py``, ``aggregate.py`` und
``validate_deliverables.py`` importiert. Es enthält ausschliesslich Mechanik:
Laden und Prüfen der JSON-Codiertabelle, Normalisierung der
Kategoriennummern, Sortierung, die drei Aggregationsebenen, das Lesen einer
geschriebenen Arbeitsmappe sowie Text- und Zitatnormalisierung.

Inhaltliche Urteile trifft dieses Modul nicht. Es klassifiziert insbesondere
keine Kategorien und leitet keine Bewertung aus Stichwörtern ab.

Beispiele in Fehlermeldungen und Doku stammen aus einem erfundenen Übungsfall
(Möbelhersteller "Firma M.", Verlagerung der Fertigung an einen zweiten
Standort).
"""

from __future__ import annotations

import json
import math
import os
import re
import unicodedata
from difflib import SequenceMatcher

BEWERTUNGEN = ("+", "-", "o", ">")

# ARGB-Farben der Arbeitsmappe (siehe references/excel-output.md)
ROT = "FF8B0000"          # Kopfzeile der Datentabellen
BLAU = "FF1F4E79"         # Kopfzeile der Methodentabelle
GRAU_HELL = "FFF2F2F2"    # Hauptkategoriezeilen
GRAU_MITTEL = "FFD9D9D9"  # Gesamtzeile der Auszählung
WEISS = "FFFFFFFF"

SHEETS = ("Aussagen", "Kategoriensystem", "Auszählung", "Info")
HEADER_AUSSAGEN = ["Kennung", "Kategorie", "Bewertung", "Kernaussage", "Äußerung"]
HEADER_KATSYS = ["Hauptkategorie", "Unterkategorie", "Definition",
                 "Ankerbeispiel", "Codierregel"]
HEADER_AUSZAEHLUNG_FIX = ["Kategorie", "Bezeichnung", None, "+", "-", "o", ">",
                          "Gesamt Äußerungen"]

RE_KENNUNG = re.compile(r"^P\d{2}$")
RE_KATEGORIE = re.compile(r"^\d+\.\d+$")
RE_GEDANKENSTRICH = re.compile(r"[–—]")
RE_AUSLASSUNG = re.compile(r"…|\.\.\.")
RE_WEISSRAUM = re.compile(r"[\n\r\t]")


class KernFehler(Exception):
    """Fehler in Eingabedaten oder Artefakten, in deutscher Sprache formuliert."""


# ---------------------------------------------------------------------------
# Kategoriennummern
# ---------------------------------------------------------------------------

def norm_kat(nummer) -> str:
    """Normalisiert eine Kategoriennummer auf die Schreibweise ohne Schlusspunkt.

    Das Sheet ``Kategoriensystem`` führt Nummern mit Schlusspunkt (``1.1.``),
    die Sheets ``Aussagen`` und ``Auszählung`` ohne (``1.1``). Beim Verknuepfen
    wird deshalb immer normalisiert.
    """
    return str(nummer).strip().rstrip(".")


def kat_sortkey(nummer):
    """Numerischer Sortierschlüssel: ``1.10`` steht hinter ``1.9``, nicht davor."""
    teile = norm_kat(nummer).split(".")
    schluessel = []
    for teil in teile:
        try:
            schluessel.append((0, int(teil)))
        except ValueError:
            schluessel.append((1, teil))
    return tuple(schluessel)


# ---------------------------------------------------------------------------
# Textnormalisierung
# ---------------------------------------------------------------------------

_ANFUEHRUNG = {
    "„": '"', "“": '"', "”": '"', "‟": '"',
    "«": '"', "»": '"', "‹": "'", "›": "'",
    "‘": "'", "’": "'", "‚": "'",
}


def normtext(text: str) -> str:
    """Vereinheitlicht Anführungszeichen, Apostrophe und Weißraum."""
    if text is None:
        return ""
    text = unicodedata.normalize("NFC", str(text))
    for zeichen, ersatz in _ANFUEHRUNG.items():
        text = text.replace(zeichen, ersatz)
    text = text.replace(" ", " ").replace(" ", " ")
    text = RE_WEISSRAUM.sub(" ", text)
    return re.sub(r"\s+", " ", text).strip()


def ohne_klammern(text: str) -> str:
    """Entfernt eckige Klammern samt Inhalt (Anonymisierungs- und Rahmungszusaetze)."""
    return re.sub(r"\s+", " ", re.sub(r"\[[^\]]*\]", " ", text)).strip()


def klammern_entkleidet(text: str) -> str:
    """Entfernt nur die Klammerzeichen, behält den Klammerinhalt."""
    return re.sub(r"\s+", " ", text.replace("[", "").replace("]", "")).strip()


def nur_wortkern(text: str) -> str:
    """Reduziert auf kleingeschriebene Wortfolge ohne Satzzeichen."""
    text = normtext(text).lower()
    text = re.sub(r"[^\w\säöüß]", " ", text, flags=re.UNICODE)
    return re.sub(r"\s+", " ", text).strip()


# Funktionswörter tragen keinen Inhalt. Ihr Wegfall ist Glättung, ihr Auftauchen
# keine Substitution. Alles andere gilt im Zitat-Gate als inhaltstragend.
FUNKTIONSWOERTER = frozenset("""
aber alle allem allen aller alles als also am an andere anderen auch auf aus
bei beim bin bis bist bloß da dabei dadurch dafür damit dann daran darauf
das dass dem den denen denn der deren des dessen die dies diese diesem diesen
dieser dieses doch dort du durch ein eine einem einen einer eines er es etwas
euch für gegen gewesen gibt hab habe haben hat hatte hatten her hier hin
hinter ich ihm ihn ihnen ihr ihre ihrem ihren ihrer ihres im in ins ist ja
jede jedem jeden jeder jedes jetzt kann keine man mehr mein meine mit nach
nicht nichts noch nun nur ob obwohl oder ohne schon sehr sein seine seinem
seinen seiner sich sie sind so solche soll sollen sondern sonst über um und
uns unser unsere unter viel vom von vor war waren warum was weil weit weiter
welche wenn wer werde werden wie wieder wir wird wirst wo wollen worden
würde zu zum zur zwar zwischen
""".split())


def inhaltswoerter(text: str) -> set[str]:
    """Inhaltstragende Wörter eines Textes, Grundlage des Substitutionsvergleichs.

    Funktionswörter und sehr kurze Wörter bleiben außen vor: Ihr Wegfall ist
    erlaubte Glättung. Taucht dagegen ein inhaltstragendes Wort im Zitat auf,
    das in der Ursprungszeile fehlt, liegt eine Substitution vor.
    """
    return {wort for wort in nur_wortkern(text).split()
            if len(wort) > 2 and wort not in FUNKTIONSWOERTER}


def aehnlichkeit(a: str, b: str) -> float:
    """Ähnlichkeitswert 0 bis 1 auf Basis von ``difflib``."""
    matcher = SequenceMatcher(None, a, b)
    if matcher.quick_ratio() < 0.5:
        return matcher.quick_ratio()
    return matcher.ratio()


# ---------------------------------------------------------------------------
# JSON-Codiertabelle
# ---------------------------------------------------------------------------

def lade_codierung(pfad: str) -> dict:
    """Liest die JSON-Codiertabelle und prüft ihr Schema.

    Erwartete Struktur siehe ``references/beispiel_codierung.json``.
    """
    if not os.path.isfile(pfad):
        raise KernFehler(f"Eingabedatei nicht gefunden: {pfad}")
    try:
        with open(pfad, encoding="utf-8") as datei:
            daten = json.load(datei)
    except json.JSONDecodeError as fehler:
        raise KernFehler(
            f"Die Datei {pfad} ist kein gültiges JSON: {fehler}") from fehler
    if not isinstance(daten, dict):
        raise KernFehler("Die JSON-Codiertabelle muss ein Objekt auf oberster Ebene sein.")

    for feld in ("meta", "personen", "perspektivgruppen", "kategoriensystem", "aussagen"):
        if feld not in daten:
            raise KernFehler(f"Pflichtfeld '{feld}' fehlt in der JSON-Codiertabelle.")

    _pruefe_meta(daten["meta"])
    _pruefe_personen(daten["personen"])
    _pruefe_gruppen(daten["perspektivgruppen"], daten["personen"])
    _pruefe_kategoriensystem(daten["kategoriensystem"])
    _pruefe_aussagen(daten["aussagen"], daten["personen"], daten["kategoriensystem"])
    return daten


def _pruefe_meta(meta) -> None:
    if not isinstance(meta, dict):
        raise KernFehler("'meta' muss ein Objekt sein.")
    pflicht = ("projekt", "gegenstand", "stand", "umfang", "auswertungsmethode",
               "codierregeln", "kernaussagen", "anonymisierungshinweis",
               "n_personen", "version")
    fehlend = [feld for feld in pflicht if feld not in meta]
    if fehlend:
        raise KernFehler("In 'meta' fehlen die Pflichtfelder: " + ", ".join(fehlend))
    if not isinstance(meta["n_personen"], int) or meta["n_personen"] < 1:
        raise KernFehler("'meta.n_personen' muss eine positive ganze Zahl sein.")
    if not isinstance(meta["version"], int) or meta["version"] < 1:
        raise KernFehler("'meta.version' muss eine positive ganze Zahl sein.")


def _pruefe_personen(personen) -> None:
    if not isinstance(personen, list) or not personen:
        raise KernFehler("'personen' muss eine nicht leere Liste sein.")
    gesehen = set()
    erlaubt = {"transkript", "mitschrift", "protokoll", "diktat"}
    for eintrag in personen:
        if not isinstance(eintrag, dict) or "kennung" not in eintrag:
            raise KernFehler("Jeder Eintrag in 'personen' braucht das Feld 'kennung'.")
        kennung = eintrag["kennung"]
        if not RE_KENNUNG.match(str(kennung)):
            raise KernFehler(
                f"Kennung '{kennung}' verletzt das Format Pnn, zum Beispiel P01.")
        if kennung in gesehen:
            raise KernFehler(f"Kennung '{kennung}' kommt in 'personen' mehrfach vor.")
        gesehen.add(kennung)
        quelle = eintrag.get("quellenart")
        if quelle is not None and quelle not in erlaubt:
            raise KernFehler(
                f"Unbekannte Quellenart '{quelle}' bei {kennung}. "
                "Erlaubt sind: transkript, mitschrift, protokoll, diktat.")


def _pruefe_gruppen(gruppen, personen) -> None:
    if not isinstance(gruppen, list) or not gruppen:
        raise KernFehler("'perspektivgruppen' muss eine nicht leere Liste sein.")
    alle_kennungen = {p["kennung"] for p in personen}
    verteilt: list[str] = []
    for gruppe in gruppen:
        if not isinstance(gruppe, dict) or "name" not in gruppe or "kennungen" not in gruppe:
            raise KernFehler(
                "Jede Perspektivgruppe braucht die Felder 'name' und 'kennungen'.")
        if not isinstance(gruppe["kennungen"], list) or not gruppe["kennungen"]:
            raise KernFehler(
                f"Die Perspektivgruppe '{gruppe.get('name')}' enthält keine Kennungen.")
        verteilt.extend(gruppe["kennungen"])
    doppelt = sorted({k for k in verteilt if verteilt.count(k) > 1})
    if doppelt:
        raise KernFehler(
            "Die Perspektivgruppen-Registry ist nicht partitionierend. "
            "Mehrfach zugeordnet: " + ", ".join(doppelt))
    fehlend = sorted(alle_kennungen - set(verteilt))
    if fehlend:
        raise KernFehler(
            "Diese Kennungen stehen in 'personen', aber in keiner Perspektivgruppe: "
            + ", ".join(fehlend))
    unbekannt = sorted(set(verteilt) - alle_kennungen)
    if unbekannt:
        raise KernFehler(
            "Diese Kennungen stehen in der Registry, aber nicht in 'personen': "
            + ", ".join(unbekannt))


DATENTYP_FALLE = (
    "Kategoriennummern gehören als JSON-String in Anführungszeichen. Als Zahl "
    "geschrieben kollabiert 1.10 zu 1.1, die Einheit wandert lautlos in die "
    "falsche Kategorie und beide Auszählungen werden falsch. Siehe die "
    "Datentyp-Falle in references/excel-output.md Abschnitt 2.")


def _pruefe_kategoriensystem(katsys) -> None:
    if not isinstance(katsys, list) or not katsys:
        raise KernFehler("'kategoriensystem' muss eine nicht leere Liste sein.")
    nummern = set()
    for eintrag in katsys:
        for feld in ("nummer", "ebene", "bezeichnung", "definition"):
            if feld not in eintrag:
                raise KernFehler(
                    f"Im Kategoriensystem fehlt das Feld '{feld}' bei "
                    f"'{eintrag.get('nummer', '?')}'.")
        if not isinstance(eintrag["nummer"], str):
            raise KernFehler(
                f"Im Kategoriensystem steht die Nummer {eintrag['nummer']!r} als "
                f"{type(eintrag['nummer']).__name__} statt als Zeichenkette. "
                + DATENTYP_FALLE)
        nummer = norm_kat(eintrag["nummer"])
        if nummer in nummern:
            raise KernFehler(f"Kategoriennummer '{nummer}' kommt mehrfach vor.")
        nummern.add(nummer)
        if eintrag["ebene"] not in ("haupt", "unter"):
            raise KernFehler(
                f"'ebene' bei Kategorie {nummer} muss 'haupt' oder 'unter' sein.")
        if eintrag["ebene"] == "haupt":
            if not re.match(r"^\d+$", nummer):
                raise KernFehler(
                    f"Hauptkategorie '{nummer}' muss eine einzelne Zahl sein, etwa '1'.")
            if eintrag.get("ankerbeispiel") or eintrag.get("codierregel"):
                raise KernFehler(
                    f"Hauptkategorie {nummer} darf kein Ankerbeispiel und keine "
                    "Codierregel tragen. Beide Felder bleiben leer.")
        else:
            if not RE_KATEGORIE.match(nummer):
                raise KernFehler(
                    f"Unterkategorie '{nummer}' muss dem Format x.y entsprechen.")
            for feld in ("ankerbeispiel", "codierregel"):
                if not str(eintrag.get(feld, "")).strip():
                    raise KernFehler(
                        f"Unterkategorie {nummer} braucht das Pflichtfeld '{feld}'.")
            haupt = nummer.split(".")[0]
            if haupt not in nummern:
                raise KernFehler(
                    f"Zu Unterkategorie {nummer} fehlt die vorangehende "
                    f"Hauptkategoriezeile '{haupt}'.")


def _pruefe_aussagen(aussagen, personen, katsys) -> None:
    if not isinstance(aussagen, list) or not aussagen:
        raise KernFehler("'aussagen' muss eine nicht leere Liste sein.")
    kennungen = {p["kennung"] for p in personen}
    unterkategorien = {norm_kat(k["nummer"]) for k in katsys if k["ebene"] == "unter"}
    for nummer, eintrag in enumerate(aussagen, start=1):
        for feld in ("kennung", "kategorie", "bewertung", "aeusserung"):
            if not str(eintrag.get(feld, "")).strip():
                raise KernFehler(
                    f"Codiereinheit {nummer}: Pflichtfeld '{feld}' fehlt oder ist leer.")
        if eintrag["kennung"] not in kennungen:
            raise KernFehler(
                f"Codiereinheit {nummer}: Kennung '{eintrag['kennung']}' steht "
                "nicht in 'personen'.")
        if not isinstance(eintrag["kategorie"], str):
            raise KernFehler(
                f"Codiereinheit {nummer}: Die Kategorie {eintrag['kategorie']!r} "
                f"steht als {type(eintrag['kategorie']).__name__} statt als "
                "Zeichenkette. " + DATENTYP_FALLE)
        if norm_kat(eintrag["kategorie"]) not in unterkategorien:
            raise KernFehler(
                f"Codiereinheit {nummer}: Kategorie '{eintrag['kategorie']}' ist "
                "keine Unterkategorie des Kategoriensystems.")
        if eintrag["bewertung"] not in BEWERTUNGEN:
            raise KernFehler(
                f"Codiereinheit {nummer}: Bewertung '{eintrag['bewertung']}' ist "
                "nicht zulässig. Erlaubt sind +, -, o, >.")
        if not isinstance(eintrag.get("kernaussage", False), bool):
            raise KernFehler(
                f"Codiereinheit {nummer}: 'kernaussage' muss true oder false sein.")


def zeilen_aus_json(daten: dict) -> list[dict]:
    """Erzeugt die sortierten Datenzeilen des Sheets ``Aussagen``."""
    zeilen = []
    for eintrag in daten["aussagen"]:
        zeilen.append({
            "kennung": str(eintrag["kennung"]).strip(),
            "kategorie": norm_kat(eintrag["kategorie"]),
            "bewertung": str(eintrag["bewertung"]).strip(),
            "kernaussage": "*" if eintrag.get("kernaussage") else None,
            "aeusserung": str(eintrag["aeusserung"]).strip(),
            "reihenfolge": eintrag.get("reihenfolge", 0),
        })
    return sortiere(zeilen)


def sortiere(zeilen: list[dict]) -> list[dict]:
    """Dreistufige Sortierung: Kategorie numerisch, Kennung, Gespraechsverlauf."""
    return sorted(zeilen, key=lambda z: (kat_sortkey(z["kategorie"]),
                                         z["kennung"],
                                         z.get("reihenfolge", 0)))


# ---------------------------------------------------------------------------
# Die drei Aggregationsebenen
# ---------------------------------------------------------------------------

def personen_je_kategorie(zeilen, kategorien) -> dict[str, int]:
    """Ebene 1: distinkte Kennungen je Unterkategorie."""
    ergebnis = {}
    for kategorie in kategorien:
        ergebnis[kategorie] = len({z["kennung"] for z in zeilen
                                   if z["kategorie"] == kategorie})
    return ergebnis


def aeusserungen_je_bewertung(zeilen, kategorien) -> dict[str, dict[str, int]]:
    """Ebene 2: Codiereinheiten je Bewertung je Unterkategorie (Sheet Auszählung)."""
    ergebnis = {}
    for kategorie in kategorien:
        treffer = [z for z in zeilen if z["kategorie"] == kategorie]
        werte = {b: sum(1 for z in treffer if z["bewertung"] == b) for b in BEWERTUNGEN}
        werte["gesamt"] = len(treffer)
        ergebnis[kategorie] = werte
    return ergebnis


def personen_je_bewertung(zeilen, kategorien) -> dict[str, dict[str, int]]:
    """Ebene 3: distinkte Kennungen je Bewertung je Unterkategorie (Zählzeilen)."""
    ergebnis = {}
    for kategorie in kategorien:
        treffer = [z for z in zeilen if z["kategorie"] == kategorie]
        werte = {b: len({z["kennung"] for z in treffer if z["bewertung"] == b})
                 for b in BEWERTUNGEN}
        werte["personen"] = len({z["kennung"] for z in treffer})
        ergebnis[kategorie] = werte
    return ergebnis


def zaehlzeile(werte: dict, n: int) -> str:
    """Baut die Zählzeile ``X von N Personen (a +, b -, c o, d >)``."""
    return (f"{werte['personen']} von {n} Personen "
            f"({werte['+']} +, {werte['-']} -, {werte['o']} o, {werte['>']} >)")


def zaehlzeile_plausibel(werte: dict, n: int) -> bool:
    """Prüft die Regel ``max(a, b, c, d) <= X <= N``."""
    groesste = max(werte[b] for b in BEWERTUNGEN)
    return groesste <= werte["personen"] <= n


def auszaehlung_berechnen(zeilen, katsys, n_personen) -> list[dict]:
    """Berechnet das Sheet ``Auszählung`` vollstaendig aus den Codiereinheiten."""
    unter = [k for k in katsys if k["ebene"] == "unter"]
    unter = sorted(unter, key=lambda k: kat_sortkey(k["nummer"]))
    kategorien = [norm_kat(k["nummer"]) for k in unter]
    ebene2 = aeusserungen_je_bewertung(zeilen, kategorien)
    ebene1 = personen_je_kategorie(zeilen, kategorien)

    tabelle = []
    for eintrag in unter:
        kategorie = norm_kat(eintrag["nummer"])
        werte = ebene2[kategorie]
        tabelle.append({
            "kategorie": kategorie,
            "bezeichnung": eintrag["bezeichnung"],
            "personen_text": f"{ebene1[kategorie]} von {n_personen}",
            "personen": ebene1[kategorie],
            "+": werte["+"], "-": werte["-"], "o": werte["o"], ">": werte[">"],
            "gesamt": werte["gesamt"],
        })
    gesamt_personen = len({z["kennung"] for z in zeilen})
    tabelle.append({
        "kategorie": "Gesamt",
        "bezeichnung": "Alle Kategorien (aggregiert)",
        "personen_text": f"{gesamt_personen} von {n_personen}",
        "personen": gesamt_personen,
        "+": sum(r["+"] for r in tabelle),
        "-": sum(r["-"] for r in tabelle),
        "o": sum(r["o"] for r in tabelle),
        ">": sum(r[">"] for r in tabelle),
        "gesamt": sum(r["gesamt"] for r in tabelle),
    })
    return tabelle


def auszaehlung_markdown(tabelle, n_personen) -> str:
    """Formatiert die berechnete Auszählung als Markdown-Tabelle."""
    kopf = (f"| Kategorie | Bezeichnung | Personen (n von {n_personen}) "
            "| + | - | o | > | Gesamt Äußerungen |")
    trenn = "|---|---|---|---:|---:|---:|---:|---:|"
    zeilen = [kopf, trenn]
    for reihe in tabelle:
        zeilen.append(
            f"| {reihe['kategorie']} | {reihe['bezeichnung']} | {reihe['personen_text']} "
            f"| {reihe['+']} | {reihe['-']} | {reihe['o']} | {reihe['>']} "
            f"| {reihe['gesamt']} |")
    return "\n".join(zeilen)


# ---------------------------------------------------------------------------
# Zeilenhoehen
# ---------------------------------------------------------------------------

def zeilenhoehe_aussagen(text: str, zeichen_pro_zeile: int = 132,
                         hoehe_pro_zeile: float = 12.5):
    """Näherungsformel für die Zeilenhoehe im Sheet ``Aussagen``."""
    anzahl = max(1, math.ceil(len(text) / zeichen_pro_zeile))
    return None if anzahl <= 1 else hoehe_pro_zeile * anzahl


def zeilenhoehe_katsys(zellen, zeichen_pro_zeile: int = 60,
                       hoehe_pro_zeile: int = 14, deckel: int = 140):
    """Näherungsformel für die Zeilenhoehe im Sheet ``Kategoriensystem``."""
    laenge = max((len(str(z or "")) for z in zellen), default=0)
    return min(deckel, hoehe_pro_zeile * max(1, math.ceil(laenge / zeichen_pro_zeile)))


# ---------------------------------------------------------------------------
# Arbeitsmappe lesen
# ---------------------------------------------------------------------------

def lade_arbeitsmappe(pfad: str) -> dict:
    """Liest eine geschriebene Arbeitsmappe zurueck in Datenstrukturen.

    Rückgabe: ``{"sheets", "zeilen", "katsys", "auszaehlung", "info",
    "registry", "n_personen", "formelzellen"}``.
    """
    from openpyxl import load_workbook

    if not os.path.isfile(pfad):
        raise KernFehler(f"Arbeitsmappe nicht gefunden: {pfad}")
    mappe = load_workbook(pfad, data_only=False)
    ergebnis = {"pfad": pfad, "sheets": list(mappe.sheetnames)}

    formelzellen = []
    for name in mappe.sheetnames:
        blatt = mappe[name]
        for reihe in blatt.iter_rows():
            for zelle in reihe:
                if isinstance(zelle.value, str) and zelle.value.startswith("="):
                    formelzellen.append(f"{name}!{zelle.coordinate}")
    ergebnis["formelzellen"] = formelzellen

    if "Aussagen" not in mappe.sheetnames:
        raise KernFehler("Der Arbeitsmappe fehlt das Sheet 'Aussagen'.")
    blatt = mappe["Aussagen"]
    ergebnis["header_aussagen"] = [blatt.cell(row=1, column=i).value for i in range(1, 6)]
    zeilen = []
    for nummer in range(2, blatt.max_row + 1):
        werte = [blatt.cell(row=nummer, column=i).value for i in range(1, 6)]
        if all(w is None for w in werte):
            continue
        zeilen.append({
            "excel_zeile": nummer,
            "kennung": (werte[0] or "") if isinstance(werte[0], str) else werte[0],
            "kategorie": norm_kat(werte[1]) if werte[1] is not None else None,
            "kategorie_roh": werte[1],
            "bewertung": werte[2],
            "kernaussage": werte[3],
            "aeusserung": werte[4],
            "reihenfolge": nummer,
        })
    ergebnis["zeilen"] = zeilen

    katsys = []
    if "Kategoriensystem" in mappe.sheetnames:
        blatt = mappe["Kategoriensystem"]
        ergebnis["header_katsys"] = [blatt.cell(row=1, column=i).value for i in range(1, 6)]
        for nummer in range(2, blatt.max_row + 1):
            werte = [blatt.cell(row=nummer, column=i).value for i in range(1, 6)]
            if all(w is None for w in werte):
                continue
            roh = str(werte[0] or "").strip()
            katsys.append({
                "excel_zeile": nummer,
                "nummer": norm_kat(roh),
                "nummer_roh": roh,
                "ebene": "unter" if "." in norm_kat(roh) else "haupt",
                "bezeichnung": werte[1],
                "definition": werte[2],
                "ankerbeispiel": werte[3],
                "codierregel": werte[4],
            })
    ergebnis["katsys"] = katsys

    auszaehlung = []
    kopf_c = None
    if "Auszählung" in mappe.sheetnames:
        blatt = mappe["Auszählung"]
        ergebnis["header_auszaehlung"] = [blatt.cell(row=1, column=i).value
                                          for i in range(1, 9)]
        kopf_c = blatt.cell(row=1, column=3).value
        for nummer in range(2, blatt.max_row + 1):
            werte = [blatt.cell(row=nummer, column=i).value for i in range(1, 9)]
            if all(w is None for w in werte):
                continue
            auszaehlung.append({
                "excel_zeile": nummer,
                "kategorie": str(werte[0] or "").strip(),
                "bezeichnung": werte[1],
                "personen_text": werte[2],
                "+": werte[3], "-": werte[4], "o": werte[5], ">": werte[6],
                "gesamt": werte[7],
            })
    ergebnis["auszaehlung"] = auszaehlung
    ergebnis["auszaehlung_kopf_c"] = kopf_c

    info = []
    if "Info" in mappe.sheetnames:
        blatt = mappe["Info"]
        for nummer in range(1, blatt.max_row + 1):
            info.append((blatt.cell(row=nummer, column=1).value,
                         blatt.cell(row=nummer, column=2).value))
    ergebnis["info"] = info
    ergebnis["registry"] = registry_aus_info(info)
    ergebnis["n_personen"] = n_aus_info(info, kopf_c)
    return ergebnis


def info_wert(info, label: str):
    """Liefert den Wert zu einem Label des Info-Sheets, sonst ``None``."""
    ziel = label.rstrip(":").strip().casefold()
    for spalte_a, spalte_b in info:
        if spalte_a and str(spalte_a).rstrip(":").strip().casefold() == ziel:
            return spalte_b
    return None


def registry_aus_info(info) -> dict[str, list[str]]:
    """Liest die Perspektivgruppen-Registry aus dem Info-Sheet."""
    registry: dict[str, list[str]] = {}
    muster = re.compile(r"^\s*\d+\.\s*(.+?)\s*\(\s*\d+\s+Interviews?\s*\)\s*:?\s*$")
    for spalte_a, spalte_b in info:
        if not spalte_a or not spalte_b:
            continue
        treffer = muster.match(str(spalte_a))
        if treffer:
            kennungen = [k.strip() for k in str(spalte_b).split(",") if k.strip()]
            registry[treffer.group(1)] = kennungen
    return registry


def n_aus_info(info, kopf_c) -> int | None:
    """Ermittelt N aus der Registry-Ueberschrift, ersatzweise aus dem Auszählungskopf."""
    for spalte_a, _ in info:
        if not spalte_a:
            continue
        treffer = re.match(r"^\s*Perspektivgruppen\s*\(\s*(\d+)\s+Interviews",
                           str(spalte_a))
        if treffer:
            return int(treffer.group(1))
    if kopf_c:
        treffer = re.search(r"von\s+(\d+)", str(kopf_c))
        if treffer:
            return int(treffer.group(1))
    return None


# ---------------------------------------------------------------------------
# Textlints auf Äußerungen
# ---------------------------------------------------------------------------

def lint_aeusserung(text: str) -> list[str]:
    """Prüft eine Äußerung gegen die Typografieregeln. Leere Liste heisst sauber."""
    befunde = []
    if RE_WEISSRAUM.search(text):
        befunde.append("enthält Zeilenumbruch oder Tabulator")
    if RE_GEDANKENSTRICH.search(text):
        befunde.append("enthält einen Gedankenstrich")
    if RE_AUSLASSUNG.search(text):
        befunde.append("enthält Auslassungspunkte")
    if "!" in text:
        befunde.append("enthält ein Ausrufezeichen")
    if text != text.strip():
        befunde.append("hat führende oder nachlaufende Leerzeichen")
    if "  " in text:
        befunde.append("enthält doppelte Leerzeichen")
    if text.count("[") != text.count("]"):
        befunde.append("hat unpaarige eckige Klammern")
    else:
        tiefe = 0
        for zeichen in text:
            if zeichen == "[":
                tiefe += 1
            elif zeichen == "]":
                tiefe -= 1
                if tiefe < 0:
                    befunde.append("hat eine schließende eckige Klammer ohne öffnende")
                    break
    return befunde


def klammeranteil(zeilen) -> float:
    """Anteil der Äußerungen mit mindestens einer eckigen Klammer."""
    if not zeilen:
        return 0.0
    mit = sum(1 for z in zeilen if "[" in str(z["aeusserung"] or ""))
    return mit / len(zeilen)


def sternquote(zeilen) -> float:
    """Anteil der mit ``*`` markierten Codiereinheiten."""
    if not zeilen:
        return 0.0
    return sum(1 for z in zeilen if str(z.get("kernaussage") or "").strip() == "*") / len(zeilen)


# ---------------------------------------------------------------------------
# Befunde
# ---------------------------------------------------------------------------

class Befund:
    """Ein Prueffund: blockierend oder Warnung, mit Gate, Ort und Text."""

    def __init__(self, gate: str, schwere: str, ort: str, text: str, detail: str = ""):
        self.gate = gate
        self.schwere = schwere  # "fehler" oder "warnung"
        self.ort = ort
        self.text = text
        self.detail = detail

    @property
    def blockierend(self) -> bool:
        return self.schwere == "fehler"

    def als_dict(self) -> dict:
        return {"gate": self.gate, "schwere": self.schwere, "ort": self.ort,
                "text": self.text, "detail": self.detail}

    def __repr__(self) -> str:
        return f"<Befund {self.schwere} {self.gate} {self.ort}: {self.text}>"


def bilanz(fehlerzahl: int, warnzahl: int) -> str:
    """Formuliert die Schlusszeile grammatisch korrekt."""
    verstoesse = ("1 blockierender Verstoß" if fehlerzahl == 1
                  else f"{fehlerzahl} blockierende Verstöße")
    warnungen = "1 Warnung" if warnzahl == 1 else f"{warnzahl} Warnungen"
    return f"{verstoesse}, {warnungen}"


def bericht_markdown(titel: str, befunde: list[Befund]) -> str:
    """Formatiert Befunde als Markdown-Abschnitt, nach Gate gruppiert."""
    zeilen = [f"## {titel}", ""]
    if not befunde:
        zeilen.append("Keine Befunde.")
        return "\n".join(zeilen) + "\n"
    zeilen.append("| Schwere | Gate | Ort | Befund |")
    zeilen.append("|---|---|---|---|")
    ordnung = {"fehler": 0, "warnung": 1}
    for befund in sorted(befunde, key=lambda b: (ordnung.get(b.schwere, 2), b.gate, b.ort)):
        marke = "FEHLER" if befund.blockierend else "Warnung"
        text = befund.text.replace("|", "\\|")
        if befund.detail:
            text += " " + befund.detail.replace("|", "\\|")
        ort = befund.ort.replace("|", "\\|")
        zeilen.append(f"| {marke} | {befund.gate} | {ort} | {text} |")
    return "\n".join(zeilen) + "\n"
