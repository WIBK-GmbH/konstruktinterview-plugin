#!/usr/bin/env python3
"""Schreibt die Arbeitsmappe der Konstruktinterview-Auswertung und validiert sie.

Aufruf und vollständige Formatregeln: references/excel-output.md.
"""

from __future__ import annotations

import argparse
import os
import re
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

import _kern as K  # noqa: E402

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
except ImportError:  # pragma: no cover
    sys.stderr.write(
        "FEHLER: Das Paket openpyxl fehlt. Installation: pip install openpyxl\n")
    raise SystemExit(2)


HILFE = """
Schreibt aus einer JSON-Codiertabelle die vierblättrige Arbeitsmappe
(Aussagen, Kategoriensystem, Auszählung, Info), berechnet die Auszählung
vollständig neu und prüft die geschriebene Datei anschließend gegen die
28 Punkte aus references/excel-output.md.

GRUNDSATZ
  Zahlen werden generiert, nie getippt. Die Auszählung wird aus dem Sheet
  Aussagen berechnet; ein vorhandenes Auszählungsblatt wird nie übernommen.

EINGABEFORMAT (Kurzform, Vollbeispiel: references/beispiel_codierung.json)
  {
    "meta": {
      "projekt", "gegenstand", "stand", "umfang", "auswertungsmethode",
      "codierregeln", "kernaussagen", "anonymisierungshinweis",
      "ankerbeispiele"        (optional: "real" | "prototyp" | Freitext),
      "n_personen"            (ganze Zahl),
      "version"               (ganze Zahl)
    },
    "personen":           [ {"kennung": "P01",
                             "perspektivgruppe": "Fertigung",
                             "quellenart": "transkript"} ],
    "perspektivgruppen":  [ {"name": "Fertigung",
                             "kennungen": ["P01", "P03"]} ],
    "kategoriensystem":   [ {"nummer": "1", "ebene": "haupt",
                             "bezeichnung": "...",
                             "definition": "Oberkategorie für ..."},
                            {"nummer": "1.1", "ebene": "unter",
                             "bezeichnung": "...", "definition": "...",
                             "ankerbeispiel": "„...“ (P03)",
                             "codierregel": "Hier codieren, wenn ..."} ],
    "aussagen":           [ {"kennung": "P03", "kategorie": "1.1",
                             "bewertung": "-", "kernaussage": true,
                             "aeusserung": "...", "reihenfolge": 17} ]
  }

  bewertung   genau ein Zeichen aus + - o >
  kernaussage true setzt das Selektionszeichen * (Zitatpool), sonst leer
  reihenfolge Position im Gesprächsverlauf, steuert die dritte Sortierstufe
  quellenart  transkript | mitschrift | protokoll | diktat

PLATZHALTER IN meta
  Damit auch die Metadaten generiert und nicht getippt werden, dürfen diese
  Felder Platzhalter tragen, die beim Schreiben gefüllt werden:
    umfang               {n}  {einheiten}
    auswertungsmethode   {kategorien}  {hauptkategorien}
    stand                {version}
  Stehen dort feste Zahlen, prüft die Validierung sie gegen die berechneten.

ANKERBEISPIELE, PRÄZEDENZ
  Der Modus entscheidet sich in dieser Reihenfolge:
    1. eine ausdrücklich gesetzte Option --anchor-mode gewinnt immer,
    2. sonst meta.ankerbeispiele aus der JSON, wenn dort real oder prototyp
       steht,
    3. sonst real.
  Wird ein Kategoriensystem aus einer früheren Erhebung übernommen, stammen
  seine Ankerbeispiele aus fremdem Material und stehen nicht im *-Pool dieser
  Mappe. Dann gehört entweder --anchor-mode prototyp gesetzt oder es werden
  reale Anker aus dem neuen Material eingetragen; sonst warnt die Validierung
  je Kategorie.

TEILBESTAND (--teilbestand)
  Zwischenstand einer Erhebungswelle oder eines einzelnen Interviews. Der
  Prüfpunkt 19 (jede Kategorie hat mindestens eine Codiereinheit) wird zur
  Warnung statt zum blockierenden Fehler. Ein Kategoriensystem aus einer
  größeren Erhebung hat in einem einzelnen Interview fast immer unbesetzte
  Kategorien; ohne das Flag müsste der Anwender das System beschneiden und
  damit vom Berichtssystem abweichen. Für den Vollabschluss läuft der Bau ohne
  das Flag, dann ist eine leere Kategorie wieder blockierend.

BEISPIELE
  python3 build_workbook.py --input codierung.json \\
      --output "Codierung_P01-P14_260315_v3.xlsx"

  python3 build_workbook.py --input codierung.json \\
      --output "Codierung_P01-P14_260315_v3.xlsx" \\
      --anchor-mode prototyp --denylist VERTRAULICH_klarnamen.txt --comfort

  python3 build_workbook.py --input welle2.json \\
      --output "Codierung_P07-P07_260320_v1.xlsx" --teilbestand

AUSGABE
  Die .xlsx, auf stdout die berechnete Auszählung als Markdown-Tabelle und der
  Validierungsbericht. Exit-Code 1 bei blockierenden Verstößen, 2 bei
  Bedien- oder Dateifehlern, 0 sonst (Warnungen erlauben den Abschluss).
"""

ANKER_TEXT = {
    "real": ("Reale, mit '*' markierte Codiereinheiten aus dem Sheet Aussagen, "
             "wörtlich übernommen und mit Kennung in Klammern."),
    "prototyp": ("Konstruierte prototypische Zitate im Register der Befragten. "
                 "Sie stehen nicht wörtlich im Sheet Aussagen."),
}

FONT_KOPF = dict(name="Arial", size=11, bold=True, color=K.WEISS)
DUENN = Side(style="thin")
RAHMEN = Border(left=DUENN, right=DUENN, top=DUENN, bottom=DUENN)


# ---------------------------------------------------------------------------
# Metadaten
# ---------------------------------------------------------------------------

def fuelle_meta(meta: dict, n_personen: int, einheiten: int,
                kategorien: int, hauptkategorien: int) -> dict:
    """Ersetzt Platzhalter in den Metadaten durch die berechneten Werte."""
    gefuellt = dict(meta)
    ersetzungen = {
        "umfang": {"n": n_personen, "einheiten": einheiten},
        "auswertungsmethode": {"kategorien": kategorien,
                               "hauptkategorien": hauptkategorien},
        "stand": {"version": meta.get("version")},
    }
    for feld, werte in ersetzungen.items():
        wert = str(gefuellt.get(feld, ""))
        for name, zahl in werte.items():
            wert = wert.replace("{" + name + "}", str(zahl))
        gefuellt[feld] = wert
    return gefuellt


# ---------------------------------------------------------------------------
# Sheets schreiben
# ---------------------------------------------------------------------------

def _kopfzeile(blatt, header, farbe, links_ab=None, hoehe=None, rahmen=False):
    for spalte, text in enumerate(header, start=1):
        zelle = blatt.cell(row=1, column=spalte, value=text)
        zelle.font = Font(**FONT_KOPF)
        zelle.fill = PatternFill("solid", fgColor=farbe)
        zelle.alignment = Alignment(
            horizontal="left" if links_ab and spalte in links_ab else "center",
            vertical="center", wrap_text=True)
        if rahmen:
            zelle.border = RAHMEN
    if hoehe:
        blatt.row_dimensions[1].height = hoehe


def schreibe_aussagen(mappe, zeilen, comfort: bool):
    blatt = mappe.active
    blatt.title = "Aussagen"
    _kopfzeile(blatt, K.HEADER_AUSSAGEN, K.ROT, links_ab={5})
    for spalte, breite in {"A": 12.0, "D": 14.0, "E": 110.0}.items():
        blatt.column_dimensions[spalte].width = breite

    for nummer, zeile in enumerate(zeilen, start=2):
        werte = [zeile["kennung"], zeile["kategorie"], zeile["bewertung"],
                 zeile["kernaussage"], zeile["aeusserung"]]
        for spalte, wert in enumerate(werte, start=1):
            zelle = blatt.cell(row=nummer, column=spalte, value=wert)
            zelle.font = Font(name="Arial", size=10)
            if spalte == 5:
                zelle.alignment = Alignment(horizontal="left", vertical="top",
                                            wrap_text=True)
            else:
                zelle.alignment = Alignment(horizontal="center", vertical="top")
            if spalte == 2:
                zelle.number_format = "@"
        hoehe = K.zeilenhoehe_aussagen(zeile["aeusserung"])
        if hoehe:
            blatt.row_dimensions[nummer].height = hoehe

    if comfort:
        blatt.freeze_panes = "A2"
        blatt.auto_filter.ref = f"A1:E{len(zeilen) + 1}"
    return blatt


def schreibe_kategoriensystem(mappe, katsys):
    blatt = mappe.create_sheet("Kategoriensystem")
    _kopfzeile(blatt, K.HEADER_KATSYS, K.BLAU, links_ab={1, 2, 3, 4, 5},
               rahmen=True)
    for spalte, breite in {"A": 16.0, "B": 45.0}.items():
        blatt.column_dimensions[spalte].width = breite

    geordnet = sorted(katsys, key=lambda k: K.kat_sortkey(k["nummer"]))
    for nummer, eintrag in enumerate(geordnet, start=2):
        haupt = eintrag["ebene"] == "haupt"
        werte = [K.norm_kat(eintrag["nummer"]) + ".",
                 eintrag["bezeichnung"],
                 eintrag["definition"],
                 None if haupt else eintrag.get("ankerbeispiel"),
                 None if haupt else eintrag.get("codierregel")]
        for spalte, wert in enumerate(werte, start=1):
            zelle = blatt.cell(row=nummer, column=spalte, value=wert)
            zelle.font = Font(name="Arial", size=11, bold=haupt)
            zelle.alignment = Alignment(horizontal="left", vertical="center",
                                        wrap_text=True)
            if spalte == 1:
                zelle.number_format = "@"
            if haupt:
                zelle.fill = PatternFill("solid", fgColor=K.GRAU_HELL)
        blatt.row_dimensions[nummer].height = K.zeilenhoehe_katsys(werte)
    return blatt


def schreibe_auszaehlung(mappe, tabelle, n_personen):
    blatt = mappe.create_sheet("Auszählung")
    header = list(K.HEADER_AUSZAEHLUNG_FIX)
    header[2] = f"Personen (n von {n_personen})"
    _kopfzeile(blatt, header, K.ROT, links_ab={2}, hoehe=28.0)
    for spalte, breite in {"A": 12.0, "B": 58.0, "C": 20.0, "D": 10.0,
                           "H": 18.0}.items():
        blatt.column_dimensions[spalte].width = breite

    for nummer, reihe in enumerate(tabelle, start=2):
        gesamtzeile = reihe["kategorie"] == "Gesamt"
        werte = [reihe["kategorie"], reihe["bezeichnung"], reihe["personen_text"],
                 reihe["+"], reihe["-"], reihe["o"], reihe[">"], reihe["gesamt"]]
        for spalte, wert in enumerate(werte, start=1):
            zelle = blatt.cell(row=nummer, column=spalte, value=wert)
            zelle.font = Font(name="Arial", size=10, bold=gesamtzeile)
            zelle.alignment = Alignment(
                horizontal="left" if spalte == 2 else "center",
                vertical="center", wrap_text=False)
            if spalte in (1, 3):
                zelle.number_format = "@"
            if gesamtzeile:
                zelle.fill = PatternFill("solid", fgColor=K.GRAU_MITTEL)
    return blatt


def schreibe_info(mappe, meta, gruppen, n_personen, anker_modus):
    blatt = mappe.create_sheet("Info")
    blatt.column_dimensions["A"].width = 38.0
    blatt.column_dimensions["B"].width = 75.0

    anker_wert = meta.get("ankerbeispiele")
    if not anker_wert or str(anker_wert).strip().lower() in ("real", "prototyp"):
        anker_wert = ANKER_TEXT[anker_modus]

    block1 = [
        ("Projekt:", meta["projekt"]),
        ("Gegenstand:", meta["gegenstand"]),
        ("Stand:", meta["stand"]),
        ("Umfang:", meta["umfang"]),
        ("Auswertungsmethode:", meta["auswertungsmethode"]),
        ("Codierregeln:", meta["codierregeln"]),
        ("Kernaussagen:", meta["kernaussagen"]),
        ("Ankerbeispiele:", anker_wert),
    ]
    zeilen: list[tuple] = list(block1)
    zeilen.append((None, None))
    zeilen.append((f"Perspektivgruppen ({n_personen} Interviews gesamt):", None))
    for laufend, gruppe in enumerate(gruppen, start=1):
        kennungen = sorted(gruppe["kennungen"])
        zeilen.append((f"{laufend}. {gruppe['name']} ({len(kennungen)} Interviews):",
                       ", ".join(kennungen)))
    zeilen.append((None, None))
    zeilen.append(("Wichtiger Hinweis zur Anonymisierung:",
                   meta["anonymisierungshinweis"]))

    for nummer, (label, wert) in enumerate(zeilen, start=1):
        zelle_a = blatt.cell(row=nummer, column=1, value=label)
        zelle_a.font = Font(name="Arial", size=10, bold=label is not None)
        zelle_a.alignment = Alignment(horizontal="left", vertical="top")
        zelle_b = blatt.cell(row=nummer, column=2, value=wert)
        zelle_b.font = Font(name="Arial", size=10)
        zelle_b.alignment = Alignment(horizontal="left", vertical="top",
                                      wrap_text=True)
    return blatt


# ---------------------------------------------------------------------------
# Validierung der geschriebenen Datei
# ---------------------------------------------------------------------------

def validiere(pfad: str, quelle: dict, denylist: list[str],
              anker_modus: str, teilbestand: bool = False) -> list[K.Befund]:
    """Führt die 28 Prüfungen aus references/excel-output.md aus.

    ``teilbestand`` stuft Prüfpunkt 19 (jede Kategorie besetzt) zur Warnung
    herab; das ist der Zwischenstand einer Welle oder eines Einzelinterviews.
    """
    befunde: list[K.Befund] = []

    def fehler(gate, ort, text, detail=""):
        befunde.append(K.Befund(gate, "fehler", ort, text, detail))

    def warnung(gate, ort, text, detail=""):
        befunde.append(K.Befund(gate, "warnung", ort, text, detail))

    mappe = K.lade_arbeitsmappe(pfad)

    # 1 bis 3: Struktur
    if mappe["sheets"] != list(K.SHEETS):
        fehler("struktur", "Workbook",
               "Die Sheets stehen nicht in der Reihenfolge "
               "Aussagen, Kategoriensystem, Auszählung, Info.",
               f"Gefunden: {mappe['sheets']}")
    if mappe.get("header_aussagen") != K.HEADER_AUSSAGEN:
        fehler("struktur", "Aussagen!1", "Kopfzeile weicht von der Spezifikation ab.",
               f"Gefunden: {mappe.get('header_aussagen')}")
    if mappe.get("header_katsys") != K.HEADER_KATSYS:
        fehler("struktur", "Kategoriensystem!1",
               "Kopfzeile weicht von der Spezifikation ab.",
               f"Gefunden: {mappe.get('header_katsys')}")
    kopf_soll = list(K.HEADER_AUSZAEHLUNG_FIX)
    kopf_ist = mappe.get("header_auszaehlung") or []
    if len(kopf_ist) != 8 or [kopf_ist[i] for i in (0, 1, 3, 4, 5, 6, 7)] != \
            [kopf_soll[i] for i in (0, 1, 3, 4, 5, 6, 7)]:
        fehler("struktur", "Auszählung!1", "Kopfzeile weicht von der Spezifikation ab.",
               f"Gefunden: {kopf_ist}")
    if kopf_ist and not re.match(r"^Personen \(n von \d+\)$", str(kopf_ist[2] or "")):
        fehler("struktur", "Auszählung!C1",
               "Der Kopf von Spalte C muss 'Personen (n von <N>)' lauten.",
               f"Gefunden: {kopf_ist[2]!r}")
    if mappe["formelzellen"]:
        fehler("struktur", "Workbook", "Die Mappe enthält Formelzellen.",
               "Betroffen: " + ", ".join(mappe["formelzellen"][:10]))

    zeilen = mappe["zeilen"]
    katsys = mappe["katsys"]
    unterkategorien = [k["nummer"] for k in katsys if k["ebene"] == "unter"]
    registry = mappe["registry"]
    registry_kennungen = {k for liste in registry.values() for k in liste}
    n_personen = mappe["n_personen"] or quelle["meta"]["n_personen"]

    # 4 bis 13: Felder und Sortierung im Sheet Aussagen
    texte: dict[str, int] = {}
    for zeile in zeilen:
        ort = f"Aussagen!{zeile['excel_zeile']}"
        for spalte, feld in (("A", "kennung"), ("B", "kategorie"),
                             ("C", "bewertung"), ("E", "aeusserung")):
            if zeile[feld] in (None, ""):
                fehler("felder", ort, f"Spalte {spalte} ist leer.")
        if zeile["bewertung"] not in K.BEWERTUNGEN:
            fehler("felder", ort,
                   f"Bewertung {zeile['bewertung']!r} ist nicht zulässig.",
                   "Erlaubt sind +, -, o, >.")
        kern = zeile["kernaussage"]
        if kern not in (None, "", "*"):
            fehler("felder", ort,
                   f"Spalte Kernaussage enthält {kern!r} statt '*' oder leer.")
        if not K.RE_KENNUNG.match(str(zeile["kennung"] or "")):
            fehler("felder", ort,
                   f"Kennung {zeile['kennung']!r} verletzt das Format Pnn.")
        elif registry_kennungen and zeile["kennung"] not in registry_kennungen:
            fehler("felder", ort,
                   f"Kennung {zeile['kennung']} steht nicht in der "
                   "Perspektivgruppen-Registry des Info-Sheets.")
        if not K.RE_KATEGORIE.match(str(zeile["kategorie"] or "")):
            fehler("felder", ort,
                   f"Kategorie {zeile['kategorie']!r} verletzt das Format x.y.")
        elif zeile["kategorie"] not in unterkategorien:
            fehler("felder", ort,
                   f"Kategorie {zeile['kategorie']} fehlt im Kategoriensystem.")
        if not isinstance(zeile["kategorie_roh"], str):
            fehler("felder", ort,
                   "Spalte Kategorie ist keine Textzelle. "
                   "Als Zahl geschrieben kollabiert 2.10 zu 2.1.")
        text = str(zeile["aeusserung"] or "")
        for befund in K.lint_aeusserung(text):
            fehler("typografie", ort, f"Äußerung {befund}.")
        schluessel = K.normtext(text)
        if schluessel in texte:
            fehler("dubletten", ort,
                   "Der Äußerungstext steht bereits in Zeile "
                   f"{texte[schluessel]}.")
        else:
            texte[schluessel] = zeile["excel_zeile"]

    sortiert = K.sortiere([dict(z) for z in zeilen])
    if [ (z["kategorie"], z["kennung"]) for z in sortiert ] != \
            [ (z["kategorie"], z["kennung"]) for z in zeilen ]:
        fehler("sortierung", "Aussagen",
               "Die Zeilen sind nicht nach Kategorie, Kennung, Gesprächsverlauf sortiert.")

    # 14 bis 19: Auszählung neu berechnen und gegenpruefen
    berechnet = K.auszaehlung_berechnen(
        zeilen,
        [{"nummer": k["nummer"], "ebene": k["ebene"], "bezeichnung": k["bezeichnung"]}
         for k in katsys],
        n_personen)
    geschrieben = mappe["auszaehlung"]
    if len(geschrieben) != len(berechnet):
        fehler("auszaehlung", "Auszählung",
               f"Das Blatt hat {len(geschrieben)} Zeilen, berechnet sind "
               f"{len(berechnet)} (Unterkategorien plus Gesamtzeile).")
    for soll, ist in zip(berechnet, geschrieben):
        ort = f"Auszählung!{ist['excel_zeile']}"
        if str(ist["kategorie"]) != soll["kategorie"]:
            fehler("auszaehlung", ort,
                   f"Kategorie {ist['kategorie']!r} erwartet {soll['kategorie']!r}.")
            continue
        if str(ist["bezeichnung"] or "") != str(soll["bezeichnung"] or ""):
            fehler("auszaehlung", ort,
                   "Bezeichnung weicht vom Kategoriensystem ab.",
                   f"Blatt: {ist['bezeichnung']!r}, System: {soll['bezeichnung']!r}")
        if str(ist["personen_text"] or "") != soll["personen_text"]:
            fehler("auszaehlung", ort,
                   f"Personenzahl {ist['personen_text']!r} statt "
                   f"{soll['personen_text']!r}.")
        for spalte in ("+", "-", "o", ">", "gesamt"):
            if ist[spalte] != soll[spalte]:
                fehler("auszaehlung", ort,
                       f"Spalte {spalte}: {ist[spalte]!r} statt {soll[spalte]!r} "
                       "(aus dem Sheet Aussagen neu berechnet).")
        summe = sum(ist[s] or 0 for s in ("+", "-", "o", ">"))
        if (ist["gesamt"] or 0) != summe:
            fehler("auszaehlung", ort,
                   f"Gesamt Äußerungen {ist['gesamt']!r} ist nicht die Summe "
                   f"von + - o > ({summe}).")
    if geschrieben:
        kategoriezeilen = [z for z in geschrieben if str(z["kategorie"]) != "Gesamt"]
        gesamtzeilen = [z for z in geschrieben if str(z["kategorie"]) == "Gesamt"]
        if not gesamtzeilen:
            fehler("auszaehlung", "Auszählung", "Die Gesamtzeile fehlt.")
        else:
            summe_h = sum(z["gesamt"] or 0 for z in kategoriezeilen)
            if summe_h != (gesamtzeilen[-1]["gesamt"] or 0) or summe_h != len(zeilen):
                fehler("auszaehlung", "Auszählung",
                       "Summe der Kategorienzeilen, Gesamtzeile und Zeilenzahl in "
                       "Aussagen stimmen nicht überein.",
                       f"{summe_h} / {gesamtzeilen[-1]['gesamt']} / {len(zeilen)}")
            gesamt_distinct = len({z["kennung"] for z in zeilen})
            if str(gesamtzeilen[-1]["personen_text"] or "") != \
                    f"{gesamt_distinct} von {n_personen}":
                fehler("auszaehlung", "Auszählung",
                       "Spalte C der Gesamtzeile ist nicht die distinkte "
                       "Gesamtzahl der Personen.")
    for kategorie in unterkategorien:
        treffer = [z for z in geschrieben if str(z["kategorie"]) == kategorie]
        if len(treffer) != 1:
            fehler("auszaehlung", "Auszählung",
                   f"Unterkategorie {kategorie} hat {len(treffer)} Zeilen statt genau einer.")
        if not any(z["kategorie"] == kategorie for z in zeilen):
            if teilbestand:
                warnung("kategorien", "Aussagen",
                        f"Kategorie {kategorie} hat keine einzige Codiereinheit.",
                        "Teilbestand: Das ist im Zwischenstand einer Welle oder "
                        "eines Einzelinterviews erwartbar. Für den Vollabschluss "
                        "muss jede Kategorie besetzt sein.")
            else:
                fehler("kategorien", "Aussagen",
                       f"Kategorie {kategorie} hat keine einzige Codiereinheit.",
                       "Ist dies ein Zwischenstand einer Welle oder ein einzeln "
                       "ausgewertetes Interview, baue mit --teilbestand.")

    # 20 bis 22: Info, Registry, Version
    umfang = str(K.info_wert(mappe["info"], "Umfang") or "")
    zahlen = {int(t.replace(".", "")) for t in re.findall(r"\d[\d.]*", umfang)}
    if len(zeilen) not in zahlen:
        fehler("info", "Info!Umfang",
               f"Die Zeile nennt nicht die berechnete Einheitenzahl {len(zeilen)}.",
               f"Gefunden: {umfang!r}")
    if n_personen not in zahlen:
        warnung("info", "Info!Umfang",
                f"Die Zeile nennt nicht die Stichprobengröße N = {n_personen}.")
    if not K.info_wert(mappe["info"], "Kernaussagen"):
        fehler("info", "Info", "Die Pflichtzeile 'Kernaussagen:' fehlt.")
    if not K.info_wert(mappe["info"], "Wichtiger Hinweis zur Anonymisierung"):
        fehler("info", "Info", "Der Anonymisierungshinweis fehlt.")
    if anker_modus == "prototyp" and not K.info_wert(mappe["info"], "Ankerbeispiele"):
        fehler("info", "Info",
               "Bei konstruierten Ankerbeispielen ist die Zeile 'Ankerbeispiele:' Pflicht.")
    if not registry:
        fehler("info", "Info", "Die Perspektivgruppen-Registry fehlt.")
    else:
        alle = [k for liste in registry.values() for k in liste]
        doppelt = sorted({k for k in alle if alle.count(k) > 1})
        if doppelt:
            fehler("info", "Info",
                   "Die Registry ist nicht partitionierend, mehrfach: " + ", ".join(doppelt))
        if len(alle) != n_personen:
            fehler("info", "Info",
                   f"Die Registry führt {len(alle)} Kennungen, N ist {n_personen}.")
        kopf_n = None
        treffer = re.search(r"von\s+(\d+)", str(mappe["auszaehlung_kopf_c"] or ""))
        if treffer:
            kopf_n = int(treffer.group(1))
        if kopf_n is not None and kopf_n != n_personen:
            fehler("info", "Auszählung!C1",
                   f"Der Kopf nennt N = {kopf_n}, die Registry N = {n_personen}.")
    stand = str(K.info_wert(mappe["info"], "Stand") or "")
    datei_version = re.search(r"_v(\d+)\.xlsx$", os.path.basename(pfad))
    info_version = re.search(r"Version\s+(\d+)", stand)
    if not datei_version:
        warnung("version", os.path.basename(pfad),
                "Der Dateiname folgt nicht dem Muster "
                "Codierung_<ErsteKennung>-<LetzteKennung>_<JJMMTT>_v<N>.xlsx.")
    if not info_version:
        warnung("version", "Info!Stand", "Die Zeile nennt keine Versionsangabe.")
    if datei_version and info_version and \
            datei_version.group(1) != info_version.group(1):
        fehler("version", "Info!Stand",
               f"Version im Dateinamen ({datei_version.group(1)}) und im Info-Sheet "
               f"({info_version.group(1)}) laufen auseinander.")

    # 23: Anonymisierung
    if denylist:
        felder = [(f"Aussagen!{z['excel_zeile']}", str(z["aeusserung"] or ""))
                  for z in zeilen]
        felder += [(f"Kategoriensystem!{k['excel_zeile']}", str(k["ankerbeispiel"] or ""))
                   for k in katsys]
        for ort, text in felder:
            unten = text.casefold()
            for begriff in denylist:
                if begriff and begriff.casefold() in unten:
                    fehler("anonymisierung", ort,
                           f"Denylist-Treffer: {begriff!r}.")
    else:
        warnung("anonymisierung", "Aufruf", "Denylist-Gate nicht geprüft.",
                "Ohne --denylist entfällt Prüfpunkt 23, der Abgleich der Spalten "
                "Äußerung und Ankerbeispiel gegen die Klarnamensliste. Lege die "
                "Denylist aus dem Pseudonymschlüssel an und übergib sie.")

    # 24 bis 28: Warnungen
    quote = K.sternquote(zeilen)
    if not 0.15 <= quote <= 0.30:
        warnung("plausibilitaet", "Aussagen",
                f"Die *-Quote liegt bei {quote:.1%}, Zielkorridor 15 bis 30 Prozent.")

    # 28: *-Quote je Person. Die Gesamtquote kann im Korridor liegen, während
    # einzelne Personen den Zitatpool tragen oder gar nicht darin vorkommen.
    je_person: dict[str, list] = {}
    for zeile in zeilen:
        je_person.setdefault(str(zeile["kennung"]), []).append(zeile)
    for kennung in sorted(je_person):
        eigene = je_person[kennung]
        sterne = sum(1 for z in eigene if str(z["kernaussage"] or "") == "*")
        anteil = sterne / len(eigene)
        if anteil > 0.45:
            warnung("plausibilitaet", f"Kennung {kennung}",
                    f"{sterne} von {len(eigene)} Einheiten sind mit * markiert "
                    f"({anteil:.1%}, Schwelle 45 Prozent).",
                    "Diese Person trägt einen übergroßen Teil des Zitatpools. "
                    "Prüfe, ob die Auswahl auch nüchterne Aussagen anderer "
                    "Personen abbildet.")
        elif len(eigene) >= 10 and anteil < 0.05:
            warnung("plausibilitaet", f"Kennung {kennung}",
                    f"Nur {sterne} von {len(eigene)} Einheiten sind mit * "
                    f"markiert ({anteil:.1%}, Schwelle 5 Prozent).",
                    "Diese Person ist im Ergebnisdokument kaum belegbar, "
                    "obwohl sie ausreichend Material beigetragen hat.")

    for kategorie in unterkategorien:
        treffer = [z for z in zeilen if z["kategorie"] == kategorie]
        if treffer and not any(str(z["kernaussage"] or "") == "*" for z in treffer):
            warnung("plausibilitaet", f"Kategorie {kategorie}",
                    "Keine einzige Äußerung ist mit * markiert; die Kategorie ist "
                    "im Ergebnisdokument nicht belegbar.")
        kennungen = {z["kennung"] for z in treffer}
        if 0 < len(kennungen) < 3:
            warnung("plausibilitaet", f"Kategorie {kategorie}",
                    f"Nur {len(kennungen)} verschiedene Kennungen. Prüfe, ob die "
                    "Kategorie eine Einzelmeinung abbildet.")
    if zeilen:
        anteil_o = sum(1 for z in zeilen if z["bewertung"] == "o") / len(zeilen)
        anteil_wunsch = sum(1 for z in zeilen if z["bewertung"] == ">") / len(zeilen)
        if anteil_o > 0.40:
            warnung("plausibilitaet", "Aussagen",
                    f"Der Anteil neutraler Einheiten liegt bei {anteil_o:.1%}. "
                    "Die Codierung ist vermutlich zu deskriptiv.")
        if anteil_wunsch < 0.15:
            warnung("plausibilitaet", "Aussagen",
                    f"Der Anteil der Veränderungswünsche liegt bei {anteil_wunsch:.1%}. "
                    "Prüfe, ob die Vorrangregel > vor - angewandt wurde.")
    anteil_klammern = K.klammeranteil(zeilen)
    if anteil_klammern > 0.40:
        warnung("plausibilitaet", "Aussagen",
                f"{anteil_klammern:.1%} der Äußerungen tragen eckige Klammern. "
                "Das deutet auf zu kontextabhängige Codiereinheiten.")

    # Ankerbeispiele
    sternpool = {K.normtext(z["aeusserung"]) for z in zeilen
                 if str(z["kernaussage"] or "") == "*"}
    for eintrag in katsys:
        if eintrag["ebene"] != "unter":
            continue
        anker = str(eintrag["ankerbeispiel"] or "")
        ort = f"Kategoriensystem!{eintrag['excel_zeile']}"
        if "[" in anker or "]" in anker:
            warnung("anker", ort,
                    "Das Ankerbeispiel enthält eckige Klammern. Wähle ein "
                    "Ankerbeispiel, das ohne Anonymisierungseingriff auskommt.")
        if anker_modus == "real":
            kern = re.sub(r"\s*\(P\d{2}\)\s*$", "", anker).strip()
            kern = K.normtext(kern).strip('"').strip()
            if kern and kern not in sternpool:
                warnung("anker", ort,
                        "Das Ankerbeispiel steht nicht wörtlich als *-markierte "
                        "Codiereinheit im Sheet Aussagen (Modus real).")
            if not re.search(r"\(P\d{2}\)\s*$", anker):
                warnung("anker", ort,
                        "Im Modus real trägt das Ankerbeispiel die Kennung in "
                        "Klammern, etwa „...“ (P03).")
    return befunde


# ---------------------------------------------------------------------------
# Hauptprogramm
# ---------------------------------------------------------------------------

def lade_denylist(pfad: str | None) -> list[str]:
    if not pfad:
        return []
    if not os.path.isfile(pfad):
        raise K.KernFehler(f"Denylist nicht gefunden: {pfad}")
    begriffe = []
    with open(pfad, encoding="utf-8") as datei:
        for zeile in datei:
            zeile = zeile.strip()
            if zeile and not zeile.startswith("#"):
                begriffe.append(zeile)
    return begriffe


def main(argv=None) -> int:
    parser = argparse.ArgumentParser(
        prog="build_workbook.py",
        description="Schreibt und validiert die Arbeitsmappe der Auswertung.",
        epilog=HILFE,
        formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--input", required=True,
                        help="JSON-Codiertabelle (Pflicht)")
    parser.add_argument("--output", required=True,
                        help="Zieldatei .xlsx, Muster "
                             "Codierung_<ErsteKennung>-<LetzteKennung>_<JJMMTT>_v<N>.xlsx")
    parser.add_argument("--anchor-mode", choices=("real", "prototyp"), default=None,
                        help="Herkunft der Ankerbeispiele: reale *-markierte "
                             "Codiereinheit oder konstruierter Prototyp. "
                             "Ausdrücklich gesetzt gewinnt diese Option vor "
                             "meta.ankerbeispiele; ohne beides gilt real")
    parser.add_argument("--denylist", default=None,
                        help="Textdatei mit Klarnamen, Orten und Einheiten, "
                             "eine je Zeile; ein Treffer ist blockierend")
    parser.add_argument("--teilbestand", action="store_true",
                        help="Zwischenstand einer Welle oder eines einzelnen "
                             "Interviews: unbesetzte Kategorien sind Warnung "
                             "statt blockierendem Fehler (Prüfpunkt 19)")
    parser.add_argument("--comfort", action="store_true",
                        help="setzt zusätzlich freeze_panes und AutoFilter im "
                             "Sheet Aussagen")
    args = parser.parse_args(argv)

    try:
        daten = K.lade_codierung(args.input)
        denylist = lade_denylist(args.denylist)
    except K.KernFehler as fehler:
        sys.stderr.write(f"FEHLER: {fehler}\n")
        return 2

    # Präzedenz des Ankermodus: ausdrücklich gesetzte Option, sonst
    # meta.ankerbeispiele, sonst real. Die Option ist bewusst ohne argparse-
    # Default, damit „nicht gesetzt“ von „auf real gesetzt“ unterscheidbar
    # bleibt und die JSON nicht stillschweigend überstimmt wird.
    anker_modus = args.anchor_mode
    if anker_modus is None:
        gesetzt = str(daten["meta"].get("ankerbeispiele", "")).strip().lower()
        anker_modus = gesetzt if gesetzt in ("real", "prototyp") else "real"

    zeilen = K.zeilen_aus_json(daten)
    katsys = daten["kategoriensystem"]
    n_personen = daten["meta"]["n_personen"]
    if n_personen != len(daten["personen"]):
        sys.stderr.write(
            f"FEHLER: meta.n_personen ist {n_personen}, in 'personen' stehen "
            f"{len(daten['personen'])} Einträge.\n")
        return 2

    unter = [k for k in katsys if k["ebene"] == "unter"]
    haupt = [k for k in katsys if k["ebene"] == "haupt"]
    meta = fuelle_meta(daten["meta"], n_personen, len(zeilen), len(unter), len(haupt))
    tabelle = K.auszaehlung_berechnen(zeilen, katsys, n_personen)

    ziel = os.path.abspath(args.output)
    verzeichnis = os.path.dirname(ziel)
    if verzeichnis and not os.path.isdir(verzeichnis):
        sys.stderr.write(f"FEHLER: Zielverzeichnis existiert nicht: {verzeichnis}\n")
        return 2

    mappe = Workbook()
    schreibe_aussagen(mappe, zeilen, args.comfort)
    schreibe_kategoriensystem(mappe, katsys)
    schreibe_auszaehlung(mappe, tabelle, n_personen)
    schreibe_info(mappe, meta, daten["perspektivgruppen"], n_personen, anker_modus)
    mappe.active = 0
    try:
        mappe.save(ziel)
    except OSError as fehler:
        sys.stderr.write(f"FEHLER: Die Datei konnte nicht geschrieben werden: {fehler}\n")
        return 2

    print(f"# Arbeitsmappe geschrieben\n\n`{ziel}`\n")
    print(f"{len(zeilen)} Codiereinheiten, {len(unter)} Unterkategorien in "
          f"{len(haupt)} Hauptkategorien, N = {n_personen}.\n")
    if args.teilbestand:
        print("Modus Teilbestand: Prüfpunkt 19 (jede Kategorie besetzt) läuft als "
              "Warnung. Der Vollabschluss wird ohne --teilbestand gebaut.\n")
    print("## Berechnete Auszählung (Äußerungen je Bewertung)\n")
    print(K.auszaehlung_markdown(tabelle, n_personen))
    print()

    try:
        befunde = validiere(ziel, daten, denylist, anker_modus, args.teilbestand)
    except K.KernFehler as fehler:
        sys.stderr.write(f"FEHLER bei der Validierung: {fehler}\n")
        return 2

    print(K.bericht_markdown("Validierung der geschriebenen Datei", befunde))
    fehlerzahl = sum(1 for b in befunde if b.blockierend)
    warnzahl = len(befunde) - fehlerzahl
    print(f"Ergebnis: {K.bilanz(fehlerzahl, warnzahl)}.")
    return 1 if fehlerzahl else 0


if __name__ == "__main__":
    raise SystemExit(main())
