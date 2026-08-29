#!/usr/bin/env python3
"""Prüft die Endartefakte der Auswertung gegen die Arbeitsmappe.

Modus A  nur --workbook: die Arbeitsmappe gegen sich selbst, Auszählung wird
         aus dem Sheet Aussagen neu berechnet und gegen die geschriebenen
         Werte diffiert. Die artefaktgebundenen Gates laufen nicht und werden
         im Protokoll ausdrücklich als nicht geprüft ausgewiesen.
Modus B  --workbook plus --deck oder --report: jede Zahl und jedes Zitat des
         Ergebnisdokuments gegen die Arbeitsmappe.

Die Begriffe sind in references/gates.md Abschnitt 9 und in --help definiert.
Vollständige Gate-Beschreibung: references/gates.md.
"""

from __future__ import annotations

import argparse
import json
import os
import re
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

import _kern as K  # noqa: E402


HILFE = """
Grundsatz: Prüfen heißt neu berechnen. Ein Gate, das eine Zahl mit sich selbst
vergleicht, prüft nichts. Alle Sollwerte entstehen aus dem Sheet Aussagen.

DIE BEIDEN MODI
  Modus A  nur --workbook. Geprüft wird allein die Arbeitsmappe: Sheetfolge,
           Wertebereiche, Dubletten, die aus dem Sheet Aussagen neu berechnete
           Auszählung gegen die geschriebene, die Pflichtzeilen des Info-Sheets
           und die Registry. Mit --denylist kommt der Anonymisierungs-Lint auf
           dem Sheet Aussagen dazu, mit --key die Registry-Gegenprüfung.
           Modus A ist eine Mappenprüfung, keine Vollprüfung: Zitat-,
           Kongruenz- und Wortlaut-Gate haben ohne Endartefakt nichts zu
           prüfen. Das Protokoll führt sie deshalb ausdrücklich als nicht
           geprüft auf.
  Modus B  --workbook plus --deck oder --report. Zusätzlich läuft jede Zahl
           und jedes Zitat des Ergebnisdokuments gegen die Arbeitsmappe. Erst
           hier greifen die artefaktgebundenen Gates.

GATES
  zaehl       Zählzeilen „X von N Personen (a +, b -, c o, d >)“ gegen die
              distinkten Personen je Kategorie und Bewertung; Prüfregel
              max(a, b, c, d) <= X <= N; Unterthemen <= Personenzahl ihrer
              Kategorie; dokumentweit identische Gesamtzahl der Äußerungen;
              keine Prozentangaben der Auswertung
  zitat       jedes Zitat gegen den *-Pool der Arbeitsmappe, Status
              VERIFIED / SMOOTHING_OK / DRIFT / NOT_FOUND, dazu Herkunft aus
              dem *-Pool, Übereinstimmung des Markers mit der Bewertung,
              Zitat-Splicing, Zitierfähigkeit der Quellenart
  kongruenz   Zahl-Text-Kongruenz: Zahlwörter, Vollständigkeits- und
              Exklusivitätsaussagen im Umfeld einer Zahl
  anonym      Denylist blockierend, Kandidatenliste als Warnung
  konsistenz  Kategoriennamen zwischen Überblick und Detail, beide Instanzen
              der Überblicksfolie, abgebrochene Sätze, Typografie
  wortlaut    Wortlaut-Sperre gegen den freigegebenen Stand (--baseline)
  alle        alle Gates (Default)

ZITATSTATUS
  VERIFIED      zeichengleich in einer *-markierten Zeile gefunden
  SMOOTHING_OK  Abweichung ohne Substitution. Zwei Fälle mit verschiedener
                Folge: eine mit … markierte Kürzung oder eine reine Tilgung
                bleibt im Ergebnisdokument, die Arbeitsmappe behält die
                vollständige Äußerung; eine Wortlautredaktion (eckige Klammer,
                Übersetzung, Anonymisierung) fließt in die Mappe zurück
  DRIFT         ein inhaltstragendes Wort steht im Zitat, das in der
                Ursprungszeile fehlt: ein Begriff wurde substituiert
  NOT_FOUND     keine Ursprungszeile zuordenbar

ANFÜHRUNGSZEICHEN
  Erkannt werden „…“ mit Klammertiefe (ein Binnenzitat zerreißt das äußere
  Zitat nicht), die Mischform „…" mit gerader Schließung, Guillemets »…« und
  ‹…› sowie, wenn sonst nichts gefunden wird, durchgängig gerade
  Anführungszeichen. Hausnorm sind „…“ mit Binnenzitaten in ‚…‘; die drei
  anderen Formen werden geprüft und zusätzlich als Typografiebefund gemeldet.
  Keine Variante bleibt unsichtbar: Ein Zitat, das das Gate nicht sieht, wäre
  ein grün protokolliertes Gate ohne Prüfung.

DENYLIST
  Textdatei, ein Begriff je Zeile, Zeilen mit # sind Kommentare. Klarnamen,
  Ortsnamen, Einheiten- und Projektbezeichnungen. Ein Treffer ist blockierend.
  Ohne --denylist meldet das Protokoll „Denylist-Gate nicht geprüft“.

PSEUDONYMSCHLÜSSEL (--key)
  Entweder die JSON-Codiertabelle (Felder personen[].quellenart und
  personen[].perspektivgruppe) oder eine Markdown-Datei mit einer Tabelle,
  deren Kopfzeile die Spalten Kennung, Perspektivgruppe und Quellenart
  benennt (normiertes Format: references/codierung.md Abschnitt 2). Aus ihr
  stammen die Registry-Gegenprüfung und die Zitierfähigkeitsregel: bei
  Mitschrift, Gedächtnisprotokoll und Diktat ist wörtliches Zitieren
  unzulässig. Fehlt --key oder die Spalte Perspektivgruppe, meldet das
  Protokoll „Registry-Gegenprüfung nicht durchgeführt“.

GLOSSAR (--glossar)
  JSON: {"kategorien": {"1.1": "verbindlicher Name"},
         "verboten": {"alter Begriff": "neuer Begriff"}}

ZITAT-REGISTRY (--zitatregistry)
  JSON: {"zitate": [{"text": "<Zitat wie im Dokument>",
                     "quelle": "Aussagen!12"}]}
  Beim Bauen des Ergebnisdokuments je Zitat notiert. Mit ihr prüft das Gate
  nicht nur, ob ein Zitat irgendwo in der Mappe steht, sondern ob der
  Fuzzy-Match auf dieselbe Zeile zeigt wie die beim Bauen notierte.

BEISPIELE
  python3 validate_deliverables.py --workbook "Codierung_P01-P14_260315_v3.xlsx"

  python3 validate_deliverables.py \\
      --workbook "Codierung_P01-P14_260315_v3.xlsx" \\
      --deck "Auswertung_v4.pptx" \\
      --denylist VERTRAULICH_klarnamen.txt \\
      --key codierung.json --json pruefprotokoll.json

  python3 validate_deliverables.py --workbook mappe.xlsx --deck neu.pptx \\
      --baseline freigegeben.pptx --gate wortlaut

  python3 validate_deliverables.py --workbook mappe.xlsx --deck deck.pptx \\
      --zitatregistry zitatregistry.json --gate zitat

AUSGABE
  Markdown auf stdout, optional ein JSON-Prüfprotokoll. Exit-Code 1, sobald
  ein blockierender Verstoß vorliegt (NOT_FOUND, DRIFT, Zitat außerhalb des
  *-Pools, falscher Marker, Zitat-Splicing, Denylist-Treffer, Zahlenabweichung,
  Zählzeile außerhalb max(a, b, c, d) <= X <= N, Registry gegen
  Pseudonymschlüssel). Kandidatenlisten und Kongruenztreffer sind Warnungen und
  brechen nicht ab. Exit-Code 2 bei Bedien- oder Dateifehlern; dazu zählt ein
  artefaktgebundenes Gate (zitat, kongruenz, wortlaut) ohne --deck oder
  --report sowie --gate wortlaut ohne --baseline.

ZULÄSSIGE ABWEICHUNGEN VON DER BEZUGSGRÖSSE
  Das Zähl-Gate erwartet dokumentweit dieselbe Bezugsgröße N. Zwei Ausnahmen
  sind deklariert und werden als Warnung statt als Fehler geführt: eine
  Bezugsgröße, die der Größe einer im Umfeld genannten Perspektivgruppe
  entspricht („3 von 3“ unter der Überschrift der Gruppe), und Zahlen eines
  früheren Standes in einem Delta-Kontext, erkennbar am Wort „Delta“ im
  Folientitel oder in der Abschnittsüberschrift.
"""

GATES = ("zaehl", "zitat", "kongruenz", "anonym", "konsistenz", "wortlaut", "alle")

RE_ZAEHLZEILE = re.compile(
    r"(\d+)\s+von\s+(\d+)\s*(Personen)?\s*\(\s*(\d+)\s*\+\s*,\s*(\d+)\s*-\s*,"
    r"\s*(\d+)\s*o\s*,\s*(\d+)\s*>\s*\)")
RE_VON = re.compile(r"(\d+)\s+von\s+(\d+)")
RE_KATNUMMER = re.compile(r"\b\d+\.\d{1,2}\b")
RE_ZITAT_GUILLEMET = re.compile(r"»([^»«]{1,600})«|‹([^‹›]{1,600})›")
RE_ZITAT_GERADE = re.compile(r"\"([^\"]{10,600})\"")
RE_EINHEITEN = re.compile(
    r"(\d[\d.]*)\s*(?:codierte\s+)?(Äußerungen|Aussagen|Codiereinheiten)",
    re.IGNORECASE)
RE_KENNUNG_IM_TEXT = re.compile(r"\bP\d{2}\b")

ZAHLWOERTER = [
    "beide", "beiden", "alle drei", "alle vier", "zu zweit", "zu dritt",
    "ein einziger", "eine einzige", "die eine", "das paar", "einzelne stimmen",
    "eine stimme", "als einziger", "als einzige", "die einzige", "der einzige",
    "keine gegenstimme", "ausnahmslos alle", "niemand", "nahezu alle",
    "die große mehrheit", "nur eine stimme",
]

KANDIDATEN_ANONYM = [
    ("meine abteilung", "Possessivpronomen verrät die Rolle der sprechenden Person"),
    ("meiner abteilung", "Possessivpronomen verrät die Rolle der sprechenden Person"),
    ("bei uns im", "Possessivkonstruktion verortet die sprechende Person"),
    ("auf unserer ebene", "Ebenen-Deiktikon verortet die Person in der Hierarchie"),
    ("in meinem bereich", "Possessivpronomen verrät die Rolle der sprechenden Person"),
    ("unsere abteilung", "Possessivpronomen verrät die Rolle der sprechenden Person"),
]


# ---------------------------------------------------------------------------
# Textstellen aus den Artefakten
# ---------------------------------------------------------------------------

def zitate_mit_tiefe(text: str) -> list[tuple[int, str, bool]]:
    """Extrahiert Zitate in ``„…“`` mit Klammertiefe.

    Liefert je Fund ``(start, inhalt, mischform)``.

    Verbindlich ist die Konvention aus ``codierung.md``: Binnenzitate stehen in
    einfachen Anführungszeichen ``‚…‘``. Altbestände setzen sie in ``„…“``;
    deshalb zählt dieser Scanner die öffnenden Zeichen und schließt das Zitat
    erst bei Tiefe null. Ein Binnenzitat bleibt so Teil des äußeren Zitats,
    statt es zu zerreißen.

    **Mischform.** Bei offener Tiefe schließt auch ein gerades Anführungszeichen
    das Zitat. Ohne diese Regel bliebe ein Zitat der Form ``„Text"`` für das
    Zitat-Gate unsichtbar: Der Tiefenscanner fände keinen Span, die
    Guillemet-Regex griffe nicht, und die Rückfallregex für gerade
    Anführungszeichen braucht zwei gerade Zeichen. Ein erfundenes oder
    gedriftetes Zitat passierte das Gate stumm. Das Flag ``mischform`` trägt den
    Fund als Typografiebefund weiter; geprüft wird er wie jedes andere Zitat.
    """
    ergebnis: list[tuple[int, str, bool]] = []
    tiefe = 0
    start = None
    mischform = False
    for index, zeichen in enumerate(text):
        if zeichen == "„":
            if tiefe == 0:
                start = index + 1
                mischform = False
            tiefe += 1
        elif zeichen in "“”\"" and tiefe > 0:
            if zeichen == "\"":
                mischform = True
            tiefe -= 1
            if tiefe == 0 and start is not None:
                inhalt = text[start:index]
                if 1 <= len(inhalt) <= 600:
                    ergebnis.append((start, inhalt, mischform))
                start = None
    return ergebnis


def ohne_zitate(text: str) -> str:
    """Blendet Zitate aller Anführungsarten aus, für Prüfungen außerhalb von Zitaten.

    Ausgeblendet wird über die Fundstelle, nicht über den Text: Ein kurzer
    Zitatinhalt, der zufällig auch außerhalb des Zitats vorkommt, würde sonst
    dort ebenfalls verschwinden und einen Befund verdecken.
    """
    spannen = [(start, start + len(inhalt))
               for start, inhalt, _ in zitate_mit_tiefe(text)]
    for treffer in RE_ZITAT_GUILLEMET.finditer(text):
        gruppe = 1 if treffer.group(1) is not None else 2
        spannen.append((treffer.start(gruppe), treffer.end(gruppe)))
    zeichen = list(text)
    for start, ende in spannen:
        for index in range(start, ende):
            zeichen[index] = " "
    return "".join(zeichen)


class Stelle:
    """Ein Absatz eines Artefakts samt Form- und Seitenkontext."""

    def __init__(self, artefakt, seite, ort, absatz, form_text, seite_text,
                 letzter_der_form=False):
        self.artefakt = artefakt
        self.seite = seite
        self.ort = ort
        self.absatz = absatz
        self.form_text = form_text
        self.seite_text = seite_text
        self.letzter_der_form = letzter_der_form


def _formtexte_aus_shape(shape, sammler):
    """Sammelt Absatzlisten je Form, rekursiv über Gruppen und Tabellen."""
    try:
        from pptx.enum.shapes import MSO_SHAPE_TYPE
        ist_gruppe = shape.shape_type == MSO_SHAPE_TYPE.GROUP
    except Exception:
        ist_gruppe = False
    if ist_gruppe:
        for unterform in shape.shapes:
            _formtexte_aus_shape(unterform, sammler)
        return
    if getattr(shape, "has_table", False):
        for reihe in shape.table.rows:
            for zelle in reihe.cells:
                absaetze = [p.text for p in zelle.text_frame.paragraphs]
                if any(a.strip() for a in absaetze):
                    sammler.append(absaetze)
        return
    if getattr(shape, "has_text_frame", False):
        absaetze = [p.text for p in shape.text_frame.paragraphs]
        if any(a.strip() for a in absaetze):
            sammler.append(absaetze)


def stellen_aus_pptx(pfad: str, artefakt: str = "Deck") -> list[Stelle]:
    """Extrahiert alle Folientexte, inklusive Gruppen und Tabellenzellen."""
    try:
        from pptx import Presentation
    except ImportError:
        raise K.KernFehler(
            "Das Paket python-pptx fehlt. Installation: pip install python-pptx")
    if not os.path.isfile(pfad):
        raise K.KernFehler(f"Ergebnisdokument nicht gefunden: {pfad}")
    praesentation = Presentation(pfad)
    stellen: list[Stelle] = []
    for nummer, folie in enumerate(praesentation.slides, start=1):
        formen: list[list[str]] = []
        for shape in folie.shapes:
            _formtexte_aus_shape(shape, formen)
        seite_text = "\n".join("\n".join(f) for f in formen)
        for absaetze in formen:
            form_text = "\n".join(absaetze)
            echte = [i for i, a in enumerate(absaetze) if a.strip()]
            letzter = echte[-1] if echte else None
            for index, absatz in enumerate(absaetze):
                if not absatz.strip():
                    continue
                stellen.append(Stelle(artefakt, nummer, f"Folie {nummer}",
                                      absatz, form_text, seite_text,
                                      index == letzter))
    return stellen


def stellen_aus_markdown(pfad: str, artefakt: str = "Bericht") -> list[Stelle]:
    """Zerlegt eine Markdown-Datei in Absätze; jede Zeile ist eine Stelle.

    Der Formkontext einer Zeile ist ihre letzte vorangehende Überschrift. Nur so
    erkennen die Gates, zu welchem Abschnitt eine Tabellenzeile gehört, etwa die
    Gruppenüberschrift über einer nach Perspektivgruppen aufgeschlüsselten
    Tabelle oder die Überschrift eines Delta-Abschnitts.
    """
    if not os.path.isfile(pfad):
        raise K.KernFehler(f"Bericht nicht gefunden: {pfad}")
    with open(pfad, encoding="utf-8") as datei:
        text = datei.read()
    stellen = []
    ueberschrift = ""
    for nummer, zeile in enumerate(text.splitlines(), start=1):
        if not zeile.strip():
            continue
        if re.match(r"^#{1,6}\s", zeile):
            ueberschrift = zeile.strip()
        form_text = f"{ueberschrift}\n{zeile}" if ueberschrift else zeile
        stellen.append(Stelle(artefakt, 0, f"Zeile {nummer}", zeile,
                              form_text, text, False))
    return stellen


# ---------------------------------------------------------------------------
# Kontext aus der Arbeitsmappe
# ---------------------------------------------------------------------------

def baue_kontext(pfad: str) -> dict:
    """Berechnet alle Sollwerte aus der Arbeitsmappe neu."""
    mappe = K.lade_arbeitsmappe(pfad)
    zeilen = mappe["zeilen"]
    unter = [k for k in mappe["katsys"] if k["ebene"] == "unter"]
    unter.sort(key=lambda k: K.kat_sortkey(k["nummer"]))
    nummern = [k["nummer"] for k in unter]
    n_personen = mappe["n_personen"] or len({z["kennung"] for z in zeilen})

    pool = []
    for zeile in zeilen:
        text = str(zeile["aeusserung"] or "")
        pool.append({
            "excel_zeile": zeile["excel_zeile"],
            "kennung": zeile["kennung"],
            "kategorie": zeile["kategorie"],
            "bewertung": zeile["bewertung"],
            "stern": str(zeile["kernaussage"] or "").strip() == "*",
            "text": text,
            "norm": K.normtext(text),
            "norm_ohne_klammern": K.normtext(K.ohne_klammern(text)),
            "kern": K.nur_wortkern(text),
            "kern_ohne_klammern": K.nur_wortkern(K.ohne_klammern(text)),
        })

    return {
        "mappe": mappe,
        "zeilen": zeilen,
        "pool": pool,
        "kategorien": nummern,
        "bezeichnungen": {k["nummer"]: str(k["bezeichnung"] or "") for k in unter},
        "n": n_personen,
        "einheiten": len(zeilen),
        "ebene1": K.personen_je_kategorie(zeilen, nummern),
        "ebene2": K.aeusserungen_je_bewertung(zeilen, nummern),
        "ebene3": K.personen_je_bewertung(zeilen, nummern),
        "registry": mappe["registry"],
    }


# ---------------------------------------------------------------------------
# Modus A: Gates auf der Arbeitsmappe
# ---------------------------------------------------------------------------

def gate_arbeitsmappe(kontext: dict) -> list[K.Befund]:
    befunde: list[K.Befund] = []
    mappe = kontext["mappe"]
    zeilen = kontext["zeilen"]
    kategorien = kontext["kategorien"]
    n_personen = kontext["n"]

    def fehler(ort, text, detail=""):
        befunde.append(K.Befund("mappe", "fehler", ort, text, detail))

    def warnung(ort, text, detail=""):
        befunde.append(K.Befund("mappe", "warnung", ort, text, detail))

    if mappe["sheets"] != list(K.SHEETS):
        fehler("Workbook", "Die vier Sheets fehlen oder stehen in falscher Reihenfolge.",
               f"Gefunden: {mappe['sheets']}")
    if mappe["formelzellen"]:
        fehler("Workbook", "Die Mappe enthält Formelzellen.",
               ", ".join(mappe["formelzellen"][:10]))

    # Domaenen
    for zeile in zeilen:
        ort = f"Aussagen!{zeile['excel_zeile']}"
        if zeile["bewertung"] not in K.BEWERTUNGEN:
            fehler(ort, f"Bewertung {zeile['bewertung']!r} liegt außerhalb von + - o >.")
        if str(zeile["kernaussage"] or "") not in ("", "*"):
            fehler(ort, f"Spalte Kernaussage enthält {zeile['kernaussage']!r} "
                        "statt '*' oder leer.")
        if not K.RE_KENNUNG.match(str(zeile["kennung"] or "")):
            fehler(ort, f"Kennung {zeile['kennung']!r} verletzt das Format Pnn.")
        if zeile["kategorie"] not in kategorien:
            fehler(ort, f"Kategorie {zeile['kategorie']!r} steht nicht im "
                        "Kategoriensystem.")

    # Dubletten
    gesehen: dict[str, int] = {}
    for zeile in zeilen:
        schluessel = K.normtext(str(zeile["aeusserung"] or ""))
        if schluessel in gesehen:
            fehler(f"Aussagen!{zeile['excel_zeile']}",
                   f"Dublette: derselbe Äußerungstext steht in Zeile {gesehen[schluessel]}.")
        else:
            gesehen[schluessel] = zeile["excel_zeile"]

    # Auszählung neu berechnen und diffen
    berechnet = K.auszaehlung_berechnen(
        zeilen,
        [{"nummer": k["nummer"], "ebene": k["ebene"], "bezeichnung": k["bezeichnung"]}
         for k in mappe["katsys"]],
        n_personen)
    geschrieben = mappe["auszaehlung"]
    if len(berechnet) != len(geschrieben):
        fehler("Auszählung",
               f"Das Blatt hat {len(geschrieben)} Zeilen, berechnet sind {len(berechnet)}.")
    nach_kategorie = {str(z["kategorie"]): z for z in geschrieben}
    for soll in berechnet:
        ist = nach_kategorie.get(soll["kategorie"])
        if ist is None:
            fehler("Auszählung", f"Zeile für {soll['kategorie']} fehlt.")
            continue
        ort = f"Auszählung!{ist['excel_zeile']}"
        if str(ist["personen_text"] or "") != soll["personen_text"]:
            fehler(ort, "Personenzählung weicht ab.",
                   f"Blatt {ist['personen_text']!r}, berechnet {soll['personen_text']!r}")
        for spalte in ("+", "-", "o", ">", "gesamt"):
            if ist[spalte] != soll[spalte]:
                fehler(ort, f"Spalte {spalte} weicht ab.",
                       f"Blatt {ist[spalte]!r}, berechnet {soll[spalte]!r}")
        if str(ist["bezeichnung"] or "") != str(soll["bezeichnung"] or ""):
            fehler(ort, "Bezeichnung weicht vom Kategoriensystem ab.",
                   f"Blatt {ist['bezeichnung']!r}, System {soll['bezeichnung']!r}")

    # Info
    if not K.info_wert(mappe["info"], "Kernaussagen"):
        fehler("Info", "Die Legende der *-Spalte fehlt. Ohne die Zeile "
                       "'Kernaussagen:' ist die Spalte für Dritte nicht "
                       "interpretierbar.")
    if not K.info_wert(mappe["info"], "Wichtiger Hinweis zur Anonymisierung"):
        fehler("Info", "Der Anonymisierungshinweis fehlt.")
    if not mappe["registry"]:
        fehler("Info", "Die Perspektivgruppen-Registry fehlt.")
    else:
        alle = [k for liste in mappe["registry"].values() for k in liste]
        if len(set(alle)) != len(alle):
            fehler("Info", "Die Registry ist nicht partitionierend.")
        if len(alle) != n_personen:
            fehler("Info", f"Die Registry führt {len(alle)} Kennungen, N ist {n_personen}.")
        unbekannt = sorted({z["kennung"] for z in zeilen} - set(alle))
        if unbekannt:
            fehler("Aussagen", "Kennungen ohne Eintrag in der Registry: "
                   + ", ".join(unbekannt))

    quote = K.sternquote(zeilen)
    if not 0.15 <= quote <= 0.30:
        warnung("Aussagen", f"Die *-Quote liegt bei {quote:.1%}, "
                            "Zielkorridor 15 bis 30 Prozent.")
    return befunde


# ---------------------------------------------------------------------------
# Zähl-Gate
# ---------------------------------------------------------------------------

def _kategorien_im_umfeld(stelle: Stelle, bekannte) -> list[str]:
    """Kategoriennummern im Absatz, sonst in der Form, sonst auf der Folie."""
    for text in (stelle.absatz, stelle.form_text, stelle.seite_text):
        treffer = [t for t in RE_KATNUMMER.findall(text) if t in bekannte]
        if treffer:
            # Reihenfolge erhalten, Dubletten entfernen
            gesehen = []
            for t in treffer:
                if t not in gesehen:
                    gesehen.append(t)
            return gesehen
    return []


def _umfeld(stelle: Stelle) -> str:
    """Absatz, Formkontext und, bei Folien, der Text der ganzen Folie."""
    teile = [stelle.absatz, stelle.form_text]
    if stelle.seite:
        teile.append(stelle.seite_text)
    return "\n".join(teile).casefold()


def _delta_kontext(stelle: Stelle) -> bool:
    """Erkennt eine als Delta gekennzeichnete Folie oder einen Delta-Abschnitt.

    Kennzeichen ist das Wort „Delta“ im Folientitel beziehungsweise in der
    Abschnittsüberschrift der Prozessnotiz (Regel Delta-Folie in
    ``pptx-output.md`` Abschnitt 2). Dort sind Zahlen des früheren Standes
    zwangsläufig und werden nicht gegen die Arbeitsmappe geprüft.
    """
    return "delta" in _umfeld(stelle)


def _gruppen_bezugsgroesse(stelle: Stelle, gruppen: dict, nenner: int) -> str | None:
    """Name der Perspektivgruppe, deren Größe der Nenner ist und die im Umfeld steht.

    Eine nach Perspektivgruppen aufgeschlüsselte Tabelle nennt zwangsläufig die
    Gruppengröße als Bezugsgröße („3 von 3"). Das ist keine abweichende
    Bezugsgröße, sondern eine deklarierte: Der Nenner entspricht einer
    Gruppengröße der Registry und der Gruppenname steht im Umfeld der Stelle.
    """
    umfeld = _umfeld(stelle)
    for name, kennungen in (gruppen or {}).items():
        if len(kennungen) == nenner and name.casefold() in umfeld:
            return name
    return None


def gate_zaehl(stellen: list[Stelle], kontext: dict) -> list[K.Befund]:
    befunde: list[K.Befund] = []
    bekannte = set(kontext["kategorien"])
    n_personen = kontext["n"]
    ebene1 = kontext["ebene1"]
    ebene3 = kontext["ebene3"]
    gruppen = kontext["registry"] or {}

    for stelle in stellen:
        absatz = stelle.absatz
        verbraucht = []
        delta = _delta_kontext(stelle)

        for treffer in RE_ZAEHLZEILE.finditer(absatz):
            verbraucht.append(treffer.span())
            x, n = int(treffer.group(1)), int(treffer.group(2))
            hat_personen = bool(treffer.group(3))
            a, b, c, d = (int(treffer.group(i)) for i in (4, 5, 6, 7))
            ist = {"personen": x, "+": a, "-": b, "o": c, ">": d}
            gefunden = treffer.group(0)
            gruppe = _gruppen_bezugsgroesse(stelle, gruppen, n) if n != n_personen else None

            # Delta-Kontext nimmt die Zählzeile unabhängig vom Nenner aus. Eine
            # Neucodierung ändert die Bewertungszahlen, ohne N zu ändern; die
            # alte Zählzeile einer Delta-Folie ist dann trotz gleichem N nicht
            # gegen die aktuelle Mappe zu prüfen (gates.md Abschnitt 3). Die
            # Plausibilitätsprüfung max(a, b, c, d) <= X <= N bleibt davor.
            if gruppe or delta:
                if not K.zaehlzeile_plausibel(ist, n):
                    befunde.append(K.Befund(
                        "zaehl", "fehler", stelle.ort,
                        "Die Zählzeile verletzt max(a, b, c, d) <= X <= N.", gefunden))
                elif gruppe:
                    befunde.append(K.Befund(
                        "zaehl", "warnung", stelle.ort,
                        f"Zählzeile mit der Bezugsgröße der Perspektivgruppe "
                        f"{gruppe!r}; sie wird nicht gegen die Gesamtstichprobe geprüft.",
                        gefunden))
                else:
                    befunde.append(K.Befund(
                        "zaehl", "warnung", stelle.ort,
                        "Zählzeile eines früheren Standes im Delta-Kontext; sie wird "
                        "nicht gegen die Arbeitsmappe geprüft.", gefunden))
                continue

            if not hat_personen:
                befunde.append(K.Befund(
                    "zaehl", "warnung", stelle.ort,
                    "Die Zählzeile nennt die Zähleinheit nicht.",
                    f"Format ist 'X von N Personen (a +, b -, c o, d >)': {gefunden!r}"))
            if n != n_personen:
                befunde.append(K.Befund(
                    "zaehl", "fehler", stelle.ort,
                    f"Die Zählzeile nennt N = {n}, die Arbeitsmappe {n_personen}.",
                    gefunden))
            if not K.zaehlzeile_plausibel(ist, max(n, n_personen)):
                befunde.append(K.Befund(
                    "zaehl", "fehler", stelle.ort,
                    "Die Zählzeile verletzt max(a, b, c, d) <= X <= N.", gefunden))

            kandidaten = _kategorien_im_umfeld(stelle, bekannte)
            passend = [k for k in kandidaten
                       if all(ebene3[k][s] == ist[s]
                              for s in ("personen", "+", "-", "o", ">"))]
            if passend:
                continue
            if not kandidaten:
                befunde.append(K.Befund(
                    "zaehl", "warnung", stelle.ort,
                    "Zu dieser Zählzeile ließ sich keine Kategoriennummer "
                    "zuordnen; sie wurde nicht gegengeprüft.", gefunden))
                continue
            erwartet = "; ".join(
                f"{k}: " + K.zaehlzeile(ebene3[k], n_personen) for k in kandidaten)
            befunde.append(K.Befund(
                "zaehl", "fehler", stelle.ort,
                f"Zählzeile stimmt nicht mit der Arbeitsmappe überein: {gefunden!r}.",
                f"Berechnet für {erwartet}"))

        for treffer in RE_VON.finditer(absatz):
            if any(start <= treffer.start() < ende for start, ende in verbraucht):
                continue
            x, n = int(treffer.group(1)), int(treffer.group(2))
            gefunden = treffer.group(0)
            if n != n_personen:
                gruppe = _gruppen_bezugsgroesse(stelle, gruppen, n)
                if gruppe:
                    if x > n:
                        befunde.append(K.Befund(
                            "zaehl", "fehler", stelle.ort,
                            f"{gefunden!r} nennt mehr Personen als die "
                            f"Perspektivgruppe {gruppe!r} umfasst."))
                    continue
                if delta:
                    befunde.append(K.Befund(
                        "zaehl", "warnung", stelle.ort,
                        f"Bezugsgröße {gefunden!r} eines früheren Standes im "
                        "Delta-Kontext; sie wird nicht gegen die Arbeitsmappe geprüft."))
                    continue
                befunde.append(K.Befund(
                    "zaehl", "fehler", stelle.ort,
                    f"Bezugsgröße {gefunden!r} weicht von N = {n_personen} ab.",
                    "Im ganzen Dokument gilt dieselbe Bezugsgröße. Zulässige "
                    "Ausnahmen: die Größe einer im Umfeld genannten "
                    "Perspektivgruppe und Zahlen in einem Delta-Abschnitt."))
                continue
            if x > n:
                befunde.append(K.Befund(
                    "zaehl", "fehler", stelle.ort,
                    f"{gefunden!r} nennt mehr Personen als die Stichprobe umfasst."))
                continue
            treffer_kat = [t for t in RE_KATNUMMER.findall(absatz) if t in bekannte]
            if len(treffer_kat) == 1:
                kategorie = treffer_kat[0]
                if x > ebene1[kategorie]:
                    befunde.append(K.Befund(
                        "zaehl", "fehler", stelle.ort,
                        f"{gefunden!r} übersteigt die Personenzahl der Kategorie "
                        f"{kategorie} ({ebene1[kategorie]} von {n_personen}).",
                        "Ein Unterthema kann nicht mehr Personen umfassen als "
                        "seine Kategorie."))

        for treffer in RE_EINHEITEN.finditer(absatz):
            roh = treffer.group(1).replace(".", "")
            if not roh.isdigit():
                continue
            zahl = int(roh)
            if zahl != kontext["einheiten"] and zahl != n_personen:
                if delta:
                    befunde.append(K.Befund(
                        "zaehl", "warnung", stelle.ort,
                        f"Die Zahl {treffer.group(0)!r} stammt aus einem früheren "
                        "Stand und steht im Delta-Kontext; sie wird nicht gegen die "
                        "Arbeitsmappe geprüft."))
                    continue
                befunde.append(K.Befund(
                    "zaehl", "fehler", stelle.ort,
                    f"Die genannte Zahl {treffer.group(0)!r} entspricht weder der "
                    f"Zeilenzahl im Sheet Aussagen ({kontext['einheiten']}) noch N "
                    f"({n_personen})."))

        ausserhalb = ohne_zitate(absatz)
        if re.search(r"\d\s*(%|Prozent)", ausserhalb):
            befunde.append(K.Befund(
                "zaehl", "warnung", stelle.ort,
                "Prozentangabe außerhalb eines Zitats.",
                "Befunde werden ausschließlich in Rohwerten X von N ausgewiesen. "
                "Zulässig sind nur Prozentzahlen, die die Befragten selbst nennen."))
    return befunde


# ---------------------------------------------------------------------------
# Zitat-Gate
# ---------------------------------------------------------------------------

def _bester_treffer(zitat_norm: str, zitat_kern: str, eintraege):
    """Sucht die ähnlichste Zeile; liefert Eintrag und Ähnlichkeitswert."""
    bester, bestwert = None, 0.0
    for eintrag in eintraege:
        wert = max(K.aehnlichkeit(zitat_norm, eintrag["norm"]),
                   K.aehnlichkeit(zitat_kern, eintrag["kern"]),
                   K.aehnlichkeit(zitat_kern, eintrag["kern_ohne_klammern"]))
        if wert > bestwert:
            bester, bestwert = eintrag, wert
    return bester, bestwert


def _neue_inhaltswoerter(zitat: str, eintrag) -> list[str]:
    """Inhaltstragende Wörter des Zitats, die in der Ursprungszeile fehlen.

    Klammerinhalte des Zitats bleiben außen vor: eckige Klammern sind erlaubte
    Redaktion (generalisierte Rollen- und Kontextangaben) und keine
    Substitution. Bleibt nach diesem Abzug ein neues inhaltstragendes Wort
    übrig, wurde ein Begriff ersetzt und nicht bloß getilgt.
    """
    neu = K.inhaltswoerter(K.ohne_klammern(zitat)) - K.inhaltswoerter(eintrag["text"])
    return sorted(neu)


def _ist_kuerzung(zitat: str, eintrag) -> bool:
    """Prüft, ob das Zitat eine mit Auslassungszeichen gekürzte Fassung ist.

    Jedes durch ``…`` getrennte Teilstück muss in der Ursprungszeile
    zusammenhängend und in derselben Reihenfolge vorkommen. Solche Kürzungen
    leben nur im Ergebnisdokument; die Arbeitsmappe behält die vollständige
    Äußerung.
    """
    teile = [K.nur_wortkern(t) for t in re.split(r"\s*(?:…|\.\.\.)\s*", zitat)]
    teile = [t for t in teile if t]
    if len(teile) < 2:
        return False
    for basis in (eintrag["kern"], eintrag["kern_ohne_klammern"]):
        position = 0
        for teil in teile:
            gefunden = basis.find(teil, position)
            if gefunden < 0:
                break
            position = gefunden + len(teil)
        else:
            return True
    return False


def _status(zitat: str, eintrag, wert: float) -> tuple[str, str]:
    """Liefert Status und Art der Abweichung.

    Die Art unterscheidet die beiden Fälle, die bisher denselben Rat bekamen:
    ``kuerzung`` bleibt im Ergebnisdokument und fließt nicht zurück,
    ``redaktion`` betrifft den Wortlaut und muss in die Arbeitsmappe
    zurückgeschrieben werden.
    """
    if eintrag is None:
        return "NOT_FOUND", ""
    zitat_kern = K.nur_wortkern(zitat)
    if K.normtext(zitat) == eintrag["norm"]:
        return "VERIFIED", ""
    if zitat_kern in (eintrag["kern"], eintrag["kern_ohne_klammern"]):
        return "VERIFIED", ""
    if _ist_kuerzung(zitat, eintrag):
        return "SMOOTHING_OK", "kuerzung"
    if K.nur_wortkern(K.ohne_klammern(zitat)) == eintrag["kern_ohne_klammern"]:
        return "SMOOTHING_OK", "redaktion"
    if wert >= 0.75:
        if _neue_inhaltswoerter(zitat, eintrag):
            return "DRIFT", "substitution"
        return "SMOOTHING_OK", "tilgung"
    return "NOT_FOUND", ""


def _enthält(eintrag, wortfolge: str) -> bool:
    """Prüft, ob eine Wortfolge zusammenhängend in einer Codiereinheit steht."""
    nadel = f" {wortfolge} "
    return (nadel in f" {eintrag['kern']} "
            or nadel in f" {eintrag['kern_ohne_klammern']} ")


def _splicing(zitat: str, pool) -> tuple:
    """Erkennt aus zwei Codiereinheiten zusammengesetzte Zitate.

    Ein Zitat gilt als gesplict, wenn es sich an einer Stelle so teilen lässt,
    dass beide Hälften als zusammenhängende Wortfolge in zwei verschiedenen
    Codiereinheiten stehen. Ein Anführungszeichenpaar ist eine Codiereinheit
    ist eine Person.

    An der Nahtstelle stehen häufig ein bis zwei eingeschobene Fugenwörter
    („und", „aber", „dann"). Der Schnitt überspringt deshalb bis zu zwei Wörter,
    sonst bliebe ein zusammengesetztes Zitat mit Bindewort ohne Diagnose.
    """
    worte = K.nur_wortkern(zitat).split()
    if len(worte) < 10:
        return None
    for schnitt in range(5, len(worte) - 4):
        vorn = " ".join(worte[:schnitt])
        a = next((e for e in pool if _enthält(e, vorn)), None)
        if a is None:
            continue
        for fuge in (0, 1, 2):
            if schnitt + fuge > len(worte) - 4:
                break
            hinten = " ".join(worte[schnitt + fuge:])
            b = next((e for e in pool
                      if e["excel_zeile"] != a["excel_zeile"] and _enthält(e, hinten)),
                     None)
            if b is not None:
                return (a, b)
    return None


def _registry_quelle(registry: dict, zitat: str):
    """Sucht den Registry-Eintrag zu einem Zitat; vergleicht über den Wortkern."""
    kern = K.nur_wortkern(zitat)
    for eintrag in registry.get("zitate", []):
        if K.nur_wortkern(str(eintrag.get("text", ""))) == kern:
            return eintrag.get("quelle")
    return None


def _zitate_der_stelle(stelle: Stelle):
    """Liefert (Zitattext, Marker oder None, Anführungsart).

    Vier Anführungsarten werden erkannt: die Hausnorm ``„…“`` mit Klammertiefe,
    die Mischform ``„…"`` mit gerader Schließung, Guillemets ``»…«`` und
    ``‹…›`` sowie, nur wenn sonst nichts gefunden wurde, durchgängig gerade
    Anführungszeichen. Keine Variante bleibt unsichtbar; die drei letzten
    werden zusätzlich als Typografiebefund gemeldet.
    """
    absatz = stelle.absatz
    marker = None
    kopf = re.match(r"^\s*([+\-o>])\s{1,3}(?=[„\"»‹])", absatz)
    if kopf:
        marker = kopf.group(1)
    roh = [(start, inhalt, "mischform" if misch else "deutsch")
           for start, inhalt, misch in zitate_mit_tiefe(absatz)]
    for treffer in RE_ZITAT_GUILLEMET.finditer(absatz):
        gruppe = 1 if treffer.group(1) is not None else 2
        roh.append((treffer.start(gruppe), treffer.group(gruppe), "guillemet"))
    if not roh:
        roh = [(t.start(1), t.group(1), "gerade")
               for t in RE_ZITAT_GERADE.finditer(absatz)]
    roh.sort(key=lambda eintrag: eintrag[0])
    return [(inhalt.strip(), marker if index == 0 else None, art)
            for index, (_, inhalt, art) in enumerate(roh)]


def gate_zitat(stellen: list[Stelle], kontext: dict,
               quellenarten: dict | None,
               zitatregistry: dict | None = None) -> list[K.Befund]:
    befunde: list[K.Befund] = []
    pool_alle = kontext["pool"]
    pool_stern = [e for e in pool_alle if e["stern"]]
    gruppennamen = [n for n in (kontext["registry"] or {}) if len(n) > 3]
    protokoll = []

    for stelle in stellen:
        for zitat, marker, art in _zitate_der_stelle(stelle):
            if len(zitat) < 10:
                continue
            kurz = zitat if len(zitat) <= 60 else zitat[:57] + "..."
            zitat_norm = K.normtext(zitat)
            zitat_kern = K.nur_wortkern(zitat)

            treffer_stern, wert_stern = _bester_treffer(zitat_norm, zitat_kern,
                                                        pool_stern)
            treffer_alle, wert_alle = _bester_treffer(zitat_norm, zitat_kern,
                                                      pool_alle)
            if treffer_alle is not None and wert_alle > wert_stern + 0.001:
                bester, wert, aus_pool = treffer_alle, wert_alle, False
            else:
                bester, wert, aus_pool = treffer_stern, wert_stern, True
            if bester is not None and not bester["stern"]:
                aus_pool = False
            status, abweichung = _status(zitat, bester, wert)

            zuordenbar = bester is not None and status != "NOT_FOUND"
            protokoll.append({
                "ort": stelle.ort, "zitat": zitat, "status": status,
                "abweichung": abweichung,
                "aehnlichkeit": round(wert, 3),
                "quelle": (f"Aussagen!{bester['excel_zeile']}" if zuordenbar else None),
                "kennung": bester["kennung"] if zuordenbar else None,
                "im_sternpool": bool(zuordenbar and bester["stern"]),
                "marker": marker,
                "bewertung": bester["bewertung"] if zuordenbar else None,
            })

            if art == "gerade":
                befunde.append(K.Befund(
                    "zitat", "warnung", stelle.ort,
                    "Das Zitat steht in geraden Anführungszeichen.",
                    "Verwende die deutschen typografischen Zeichen „…“."))
            elif art == "mischform":
                befunde.append(K.Befund(
                    "zitat", "warnung", stelle.ort,
                    "Das Zitat ist in der Mischform „…\" gesetzt: deutsches "
                    "öffnendes, gerades schließendes Anführungszeichen.",
                    "Hausnorm ist das schließende Zeichen “ (U+201C). Das Zitat "
                    "wurde trotzdem geprüft; korrigiere die Typografie, damit "
                    "die Paarigkeitsprüfung nicht zusätzlich anschlägt."))
            elif art == "guillemet":
                befunde.append(K.Befund(
                    "zitat", "warnung", stelle.ort,
                    "Das Zitat steht in Guillemets.",
                    "Hausnorm sind die deutschen typografischen Zeichen „…“, "
                    "Binnenzitate in ‚…‘."))

            if status == "NOT_FOUND":
                befunde.append(K.Befund(
                    "zitat", "fehler", stelle.ort,
                    f"NOT_FOUND: „{kurz}“ ist in der Arbeitsmappe nicht auffindbar.",
                    f"bester Wert {wert:.2f}. Zitat entfernen und protokollieren."))
            elif status == "DRIFT":
                befunde.append(K.Befund(
                    "zitat", "fehler", stelle.ort,
                    f"DRIFT: „{kurz}“ weicht von der Ursprungszeile ab.",
                    f"Aussagen!{bester['excel_zeile']}, Wert {wert:.2f}. Ersetzt "
                    "wurde: " + ", ".join(_neue_inhaltswoerter(zitat, bester))
                    + ". Originalbegriff wiederherstellen."))
            elif status == "SMOOTHING_OK":
                if abweichung in ("kuerzung", "tilgung"):
                    hinweis = ("Gekürzt, ohne den Wortlaut zu ändern. Die "
                               "Arbeitsmappe behält die vollständige Äußerung; "
                               "die Kürzung lebt nur im Ergebnisdokument. Kein "
                               "Rückfluss nötig.")
                else:
                    hinweis = ("Der Wortlaut ist redigiert. Redaktion in die "
                               "Arbeitsmappe zurückschreiben.")
                befunde.append(K.Befund(
                    "zitat", "warnung", stelle.ort,
                    f"SMOOTHING_OK: „{kurz}“ weicht von der Ursprungszeile ab.",
                    f"Aussagen!{bester['excel_zeile']}, Wert {wert:.2f}. " + hinweis))

            if zitatregistry is not None and zuordenbar:
                deklariert = _registry_quelle(zitatregistry, zitat)
                gefunden_quelle = f"Aussagen!{bester['excel_zeile']}"
                if deklariert is None:
                    befunde.append(K.Befund(
                        "zitat", "warnung", stelle.ort,
                        f"Für „{kurz}“ steht kein Eintrag in der Zitat-Registry.",
                        "Jedes Zitat des Ergebnisdokuments trägt dort seine "
                        "Quellzeile."))
                elif str(deklariert).strip() != gefunden_quelle:
                    befunde.append(K.Befund(
                        "zitat", "fehler", stelle.ort,
                        f"Die Zitat-Registry nennt für „{kurz}“ die Quelle "
                        f"{deklariert!r}, gefunden wurde {gefunden_quelle}.",
                        "Der Fuzzy-Match zeigt auf eine andere Zeile als die "
                        "beim Bauen notierte. Zuordnung klären."))

            if bester is not None and not aus_pool and status != "NOT_FOUND":
                befunde.append(K.Befund(
                    "zitat", "fehler", stelle.ort,
                    f"„{kurz}“ stammt aus einer Zeile ohne *-Markierung.",
                    f"Aussagen!{bester['excel_zeile']}. Ohne Markierung kein Zitat."))

            if marker and bester is not None and status != "NOT_FOUND" and \
                    marker != bester["bewertung"]:
                befunde.append(K.Befund(
                    "zitat", "fehler", stelle.ort,
                    f"Der Marker {marker!r} entspricht nicht der Bewertung "
                    f"{bester['bewertung']!r} der zitierten Zeile.",
                    f"Aussagen!{bester['excel_zeile']}"))

            if status in ("DRIFT", "NOT_FOUND"):
                paar = _splicing(zitat, pool_alle)
                if paar:
                    befunde.append(K.Befund(
                        "zitat", "fehler", stelle.ort,
                        "Zitat-Splicing: der Text setzt sich aus zwei "
                        "Codiereinheiten zusammen.",
                        f"Aussagen!{paar[0]['excel_zeile']} ({paar[0]['kennung']}) "
                        f"und Aussagen!{paar[1]['excel_zeile']} ({paar[1]['kennung']}). "
                        "Ein Anführungszeichenpaar ist eine Codiereinheit ist eine Person."))

            if RE_KENNUNG_IM_TEXT.search(zitat):
                befunde.append(K.Befund(
                    "zitat", "fehler", stelle.ort,
                    f"Das Zitat trägt eine Personenkennung: „{kurz}“."))
            for name in gruppennamen:
                if name.casefold() in zitat.casefold():
                    befunde.append(K.Befund(
                        "zitat", "fehler", stelle.ort,
                        f"Das Zitat nennt die Perspektivgruppe {name!r}."))

            if not 30 <= len(zitat) <= 400:
                befunde.append(K.Befund(
                    "zitat", "warnung", stelle.ort,
                    f"Zitatlänge {len(zitat)} Zeichen liegt außerhalb von 30 bis 400."))

            if quellenarten and bester is not None and status != "NOT_FOUND":
                art = quellenarten.get(bester["kennung"])
                if art and art != "transkript":
                    befunde.append(K.Befund(
                        "zitat", "fehler", stelle.ort,
                        f"Die Quellenart {art!r} der Kennung {bester['kennung']} "
                        "erlaubt kein wörtliches Zitat.",
                        "Nur sinngemäß und ohne Anführungszeichen im Rahmentext "
                        "wiedergeben; eine Kennzeichnung heilt die fehlende "
                        "Wortlautqualität nicht."))
    kontext["zitatprotokoll"] = protokoll
    return befunde


# ---------------------------------------------------------------------------
# Konsistenz-, Kongruenz- und Anonymisierungs-Gates
# ---------------------------------------------------------------------------

def gate_konsistenz(stellen: list[Stelle], kontext: dict,
                    glossar: dict | None) -> list[K.Befund]:
    befunde: list[K.Befund] = []
    bezeichnungen = kontext["bezeichnungen"]
    bekannte = set(bezeichnungen)

    muster = re.compile(r"^\s*(\d+\.\d{1,2})\.?[\s:]+(.{2,80})$")
    gefundene_namen: dict[str, dict[str, list[str]]] = {}
    for stelle in stellen:
        treffer = muster.match(stelle.absatz.strip())
        if not treffer or treffer.group(1) not in bekannte:
            continue
        nummer = treffer.group(1)
        name = treffer.group(2).strip()
        name = re.sub(r"\s*[→>]\s*\d+\s+von\s+\d+.*$", "", name).strip()
        name = re.sub(r"\s*\(\s*\d+\s*\+.*$", "", name).strip()
        name = name.rstrip(" .:")
        if not name or RE_VON.search(name):
            continue
        gefundene_namen.setdefault(nummer, {}).setdefault(name, []).append(stelle.ort)

    for nummer, varianten in gefundene_namen.items():
        soll = bezeichnungen[nummer].strip()
        for name, orte in varianten.items():
            if K.normtext(name).casefold() != K.normtext(soll).casefold():
                befunde.append(K.Befund(
                    "konsistenz", "warnung", ", ".join(sorted(set(orte))),
                    f"Kategorie {nummer} heißt hier {name!r}, in der Arbeitsmappe "
                    f"{soll!r}.",
                    "Kategoriennamen sind in allen Artefakten wörtlich identisch."))
        if len(varianten) > 1:
            befunde.append(K.Befund(
                "konsistenz", "fehler", ", ".join(sorted(
                    {o for orte in varianten.values() for o in orte})),
                f"Kategorie {nummer} trägt im Dokument mehrere Namen: "
                + " / ".join(sorted(varianten))))

    # Beide Instanzen der Kategorienüberblicksfolie
    seiten: dict[int, str] = {}
    for stelle in stellen:
        if stelle.seite:
            seiten[stelle.seite] = stelle.seite_text
    schwelle = max(2, len(bekannte) // 2)
    ueberblick = {nummer: text for nummer, text in seiten.items()
                  if len({t for t in RE_KATNUMMER.findall(text) if t in bekannte})
                  >= schwelle}
    if len(ueberblick) >= 2:
        nummern = sorted(ueberblick)
        referenz = K.normtext(ueberblick[nummern[0]])
        for nummer in nummern[1:]:
            if K.normtext(ueberblick[nummer]) != referenz:
                befunde.append(K.Befund(
                    "konsistenz", "warnung", f"Folie {nummern[0]} und Folie {nummer}",
                    "Die beiden Instanzen der Kategorienüberblicksfolie sind nicht "
                    "zeichengleich."))

    # Abgebrochene Saetze und Typografie
    for stelle in stellen:
        text = stelle.absatz.rstrip()
        absaetze_der_form = [a for a in stelle.form_text.splitlines() if a.strip()]
        if stelle.letzter_der_form and text.endswith(":") and \
                (len(absaetze_der_form) >= 2 or len(text) > 45):
            # Eine allein stehende kurze Rubrik wie „Stimmen aus den
            # Interviews:" ist regelkonform und wird nicht gemeldet.
            befunde.append(K.Befund(
                "konsistenz", "warnung", stelle.ort,
                "Abgebrochene Karte: der letzte Absatz endet mit einem "
                "Doppelpunkt, die angekündigte Aufzählung fehlt.",
                text[-70:]))
        elif stelle.letzter_der_form and text.endswith((",", ";", "und", "oder")):
            befunde.append(K.Befund(
                "konsistenz", "warnung", stelle.ort,
                "Der letzte Absatz der Form endet mitten im Satz.", text[-70:]))
        if "„„" in text or "““" in text:
            befunde.append(K.Befund("konsistenz", "warnung", stelle.ort,
                                    "Doppeltes Anführungszeichen."))
        if text.count("„") != text.count("“") + text.count("”"):
            befunde.append(K.Befund(
                "konsistenz", "warnung", stelle.ort,
                "Anführungszeichen sind nicht paarig gesetzt.", text[:70]))
        if text.count("»") != text.count("«") or text.count("‹") != text.count("›"):
            befunde.append(K.Befund(
                "konsistenz", "warnung", stelle.ort,
                "Guillemets sind nicht paarig gesetzt.", text[:70]))
        if text.count("‚") != text.count("‘"):
            befunde.append(K.Befund(
                "konsistenz", "warnung", stelle.ort,
                "Einfache Anführungszeichen sind nicht paarig gesetzt.",
                "Binnenzitate stehen in ‚…‘. " + text[:70]))
        # Das Zitatformat schreibt zwei Leerzeichen hinter dem Marker vor.
        ohne_marker = re.sub(r"^\s*[+\-o>]\s{1,3}(?=[„\"»‹])", "", text)
        if "  " in ohne_marker:
            befunde.append(K.Befund("konsistenz", "warnung", stelle.ort,
                                    "Doppeltes Leerzeichen.", text[:70]))
        if K.RE_GEDANKENSTRICH.search(text):
            befunde.append(K.Befund(
                "konsistenz", "warnung", stelle.ort,
                "Gedankenstrich im Text. Spannen werden als „X bis Y“ geschrieben.",
                text[:70]))

    if glossar:
        for alt, neu in (glossar.get("verboten") or {}).items():
            for stelle in stellen:
                if alt.casefold() in stelle.absatz.casefold():
                    befunde.append(K.Befund(
                        "konsistenz", "warnung", stelle.ort,
                        f"Der Begriff {alt!r} sollte laut Glossar {neu!r} heißen."))
        for nummer, name in (glossar.get("kategorien") or {}).items():
            soll = bezeichnungen.get(K.norm_kat(nummer))
            if soll and K.normtext(soll) != K.normtext(name):
                befunde.append(K.Befund(
                    "konsistenz", "warnung", "Arbeitsmappe",
                    f"Glossar und Arbeitsmappe benennen Kategorie {nummer} "
                    f"verschieden: {name!r} gegen {soll!r}."))
    return befunde


def gate_kongruenz(stellen: list[Stelle], kontext: dict) -> list[K.Befund]:
    befunde: list[K.Befund] = []
    n_personen = kontext["n"]
    for stelle in stellen:
        absatz = stelle.absatz
        unten = absatz.casefold()
        zahlen = [(int(t.group(1)), int(t.group(2))) for t in RE_VON.finditer(absatz)]
        if not zahlen:
            continue
        for wort in ZAHLWOERTER:
            if wort in unten:
                befunde.append(K.Befund(
                    "kongruenz", "warnung", stelle.ort,
                    f"Zahlwort {wort!r} steht neben einer Zahlenangabe.",
                    "Nach einem Datenupdate wird die Ziffer ersetzt und der Satz "
                    f"drumherum nicht mitgelesen: {absatz[:110]!r}"))
        for x, n in zahlen:
            if n != n_personen:
                continue
            if ("nahezu alle" in unten or "große mehrheit" in unten) and x < 0.75 * n:
                befunde.append(K.Befund(
                    "kongruenz", "warnung", stelle.ort,
                    f"Mengenwort und Rohwert passen nicht zusammen: {x} von {n}.",
                    "„nahezu alle“ und „die große Mehrheit“ setzen mindestens drei "
                    "Viertel von N voraus."))
            if ("ausnahmslos alle" in unten or "keine gegenstimme" in unten) and x < n:
                befunde.append(K.Befund(
                    "kongruenz", "warnung", stelle.ort,
                    f"Vollständigkeitsaussage neben {x} von {n}.",
                    "Prüfe, ob die Aussage nach dem Datenupdate noch stimmt."))
    return befunde


def gate_anonym(stellen: list[Stelle], kontext: dict,
                denylist: list[str]) -> list[K.Befund]:
    befunde: list[K.Befund] = []
    quellen = [(s.ort, s.absatz) for s in stellen]
    quellen += [(f"Aussagen!{z['excel_zeile']}", str(z["aeusserung"] or ""))
                for z in kontext["zeilen"]]
    quellen += [(f"Kategoriensystem!{k['excel_zeile']}", str(k["ankerbeispiel"] or ""))
                for k in kontext["mappe"]["katsys"]]
    for ort, text in quellen:
        unten = text.casefold()
        for begriff in denylist:
            if begriff and begriff.casefold() in unten:
                befunde.append(K.Befund(
                    "anonym", "fehler", ort, f"Denylist-Treffer: {begriff!r}.",
                    text[:110]))
        for kandidat, grund in KANDIDATEN_ANONYM:
            if kandidat in unten:
                befunde.append(K.Befund(
                    "anonym", "warnung", ort,
                    f"Kandidat für Re-Identifikation: {kandidat!r}.",
                    f"{grund}. Zur menschlichen Entscheidung."))
    # Der Lint läuft über alle Artefakte, also auch über die Äußerungen und
    # Ankerbeispiele der Arbeitsmappe. Eine identifizierende Beschäftigtenzahl
    # dort bleibt sonst unentdeckt, bis sie zitiert wird.
    for ort, text in quellen:
        for treffer in re.finditer(r"\b(\d{2,5})\s+(Kolleg\w*|Beschäftigt\w*|"
                                   r"Mitarbeiter\w*)", text):
            befunde.append(K.Befund(
                "anonym", "warnung", ort,
                f"Beschäftigtenzahl {treffer.group(0)!r} kann den Bereich "
                "identifizieren.", "Generalisieren, etwa „die anderen Beschäftigten“."))
    return befunde


def gate_registry(kontext: dict, key_pfad, schluesselgruppen) -> list[K.Befund]:
    """Prüft die Perspektivgruppen-Registry gegen den Pseudonymschlüssel.

    Die Registry ist eine abgeleitete Struktur. Ohne Gegenprüfung gegen die
    Primärquelle driftet sie, und der Bericht weist Gruppengrößen aus, die
    niemand mehr gegen den Schlüssel gehalten hat. Bleibt die Gegenprüfung aus,
    sagt das Protokoll das ausdrücklich, statt still zu passieren.
    """
    befunde: list[K.Befund] = []
    if not key_pfad:
        return [K.Befund(
            "konsistenz", "warnung", "Aufruf",
            "Registry-Gegenprüfung nicht durchgeführt.",
            "Ohne --key fehlt die Primärquelle. Die Perspektivgruppen-Registry "
            "des Info-Sheets wurde gegen nichts gehalten.")]
    if not schluesselgruppen:
        return [K.Befund(
            "konsistenz", "warnung", os.path.basename(str(key_pfad)),
            "Registry-Gegenprüfung nicht durchgeführt.",
            "Der Pseudonymschlüssel liefert keine Gruppenzuordnung. Normiere die "
            "Tabelle mit einer Spalte 'Perspektivgruppe' (siehe codierung.md "
            "Abschnitt 2).")]
    if not kontext["registry"]:
        return [K.Befund(
            "konsistenz", "fehler", "Info",
            "Die Perspektivgruppen-Registry fehlt, der Pseudonymschlüssel führt "
            "aber Gruppen.")]
    for name, kennungen in schluesselgruppen.items():
        registry = kontext["registry"].get(name)
        if registry is None:
            befunde.append(K.Befund(
                "konsistenz", "fehler", "Info",
                f"Die Gruppe {name!r} des Pseudonymschlüssels fehlt in der Registry."))
        elif sorted(registry) != sorted(kennungen):
            befunde.append(K.Befund(
                "konsistenz", "fehler", "Info",
                f"Die Gruppe {name!r} weicht vom Pseudonymschlüssel ab.",
                f"Registry {sorted(registry)}, Schlüssel {sorted(kennungen)}"))
    for name in kontext["registry"]:
        if name not in schluesselgruppen:
            befunde.append(K.Befund(
                "konsistenz", "fehler", "Info",
                f"Die Gruppe {name!r} der Registry steht nicht im Pseudonymschlüssel."))
    return befunde


def gate_wortlaut(stellen: list[Stelle], baseline: list[Stelle]) -> list[K.Befund]:
    """Meldet jede Prosaänderung gegenüber dem freigegebenen Stand."""
    befunde: list[K.Befund] = []

    def geruest(text: str) -> str:
        return re.sub(r"\d+", "#", K.normtext(text))

    alt: dict = {}
    for stelle in baseline:
        alt.setdefault((stelle.seite, geruest(stelle.absatz)), []).append(stelle.absatz)
        alt.setdefault((None, geruest(stelle.absatz)), []).append(stelle.absatz)
    baseline_texte = {K.normtext(s.absatz) for s in baseline}

    for stelle in stellen:
        if K.normtext(stelle.absatz) in baseline_texte:
            continue
        gerippe = geruest(stelle.absatz)
        vorher = alt.get((stelle.seite, gerippe)) or alt.get((None, gerippe))
        if vorher:
            befunde.append(K.Befund(
                "wortlaut", "warnung", stelle.ort,
                "Nur Zahlen geändert, Satz im Übrigen unverändert.",
                f"vorher: {vorher[0][:90]!r}"))
        else:
            befunde.append(K.Befund(
                "wortlaut", "warnung", stelle.ort,
                "Prosaänderung gegenüber dem freigegebenen Stand.",
                f"{stelle.absatz[:110]!r}. Nach der Freigabe ändert ein "
                "Datenupdate nur Zahlen, Tippfehler und falsche Tatsachen."))

    # Gegenrichtung: Das Entfernen eines freigegebenen Absatzes ist ebenfalls
    # eine Prosaänderung und bleibt sonst unbemerkt.
    jetzt_texte = {K.normtext(s.absatz) for s in stellen}
    jetzt_geruest = {geruest(s.absatz) for s in stellen}
    for stelle in baseline:
        if K.normtext(stelle.absatz) in jetzt_texte:
            continue
        if geruest(stelle.absatz) in jetzt_geruest:
            continue
        befunde.append(K.Befund(
            "wortlaut", "warnung", stelle.ort,
            "Absatz gegenüber dem freigegebenen Stand entfernt.",
            f"{stelle.absatz[:110]!r}. Auch das Streichen freigegebener Aussagen "
            "ist eine Prosaänderung und wird einzeln begründet."))
    return befunde


# ---------------------------------------------------------------------------
# Hilfsdateien
# ---------------------------------------------------------------------------

def lade_denylist(pfad):
    if not pfad:
        return []
    if not os.path.isfile(pfad):
        raise K.KernFehler(f"Denylist nicht gefunden: {pfad}")
    begriffe = []
    with open(pfad, encoding="utf-8") as datei:
        for zeile in datei:
            zeile = zeile.strip()
            if zeile and not zeile.startswith("#"):
                begriffe.append(zeile)
    return begriffe


QUELLENARTEN = ("transkript", "mitschrift", "protokoll", "diktat")


def _tabellenfelder(zeile: str) -> list[str] | None:
    """Zerlegt eine Markdown-Tabellenzeile; liefert None, wenn es keine ist."""
    if "|" not in zeile:
        return None
    return [feld.strip() for feld in zeile.strip().strip("|").split("|")]


def lade_schluessel(pfad):
    """Liest Quellenart und Perspektivgruppe je Kennung aus JSON oder Markdown.

    Markdown: eine Tabelle, deren Kopfzeile die Spalten `Kennung`,
    `Perspektivgruppe` und `Quellenart` benennt. Das normierte Format steht in
    references/codierung.md Abschnitt 2. Ohne Spalte `Perspektivgruppe` bleibt
    die Gruppenzuordnung leer; der Aufrufer meldet dann ausdrücklich, dass die
    Registry-Gegenprüfung nicht durchgeführt wurde.
    """
    if not pfad:
        return None, None
    if not os.path.isfile(pfad):
        raise K.KernFehler(f"Pseudonymschlüssel nicht gefunden: {pfad}")
    quellenarten: dict[str, str] = {}
    gruppen: dict[str, list[str]] = {}
    if pfad.lower().endswith(".json"):
        with open(pfad, encoding="utf-8") as datei:
            daten = json.load(datei)
        for eintrag in daten.get("personen", []):
            if eintrag.get("quellenart"):
                quellenarten[eintrag["kennung"]] = eintrag["quellenart"]
            if eintrag.get("perspektivgruppe"):
                gruppen.setdefault(eintrag["perspektivgruppe"], []).append(
                    eintrag["kennung"])
        return quellenarten, gruppen

    with open(pfad, encoding="utf-8") as datei:
        zeilen = datei.readlines()

    spalten: dict[str, int] | None = None
    for zeile in zeilen:
        felder = _tabellenfelder(zeile)
        if felder is None:
            continue
        unten = [feld.casefold() for feld in felder]
        if spalten is None:
            if "kennung" in unten:
                spalten = {name: unten.index(name)
                           for name in ("kennung", "perspektivgruppe", "quellenart")
                           if name in unten}
            continue
        if all(set(feld) <= set("-: ") for feld in felder):
            continue                       # Trennzeile unter der Kopfzeile
        index = spalten.get("kennung")
        if index is None or index >= len(felder):
            continue
        if not K.RE_KENNUNG.match(felder[index]):
            continue
        kennung = felder[index]
        index = spalten.get("quellenart")
        if index is not None and index < len(felder):
            art = felder[index].casefold()
            for erlaubt in QUELLENARTEN:
                if erlaubt in art:
                    quellenarten[kennung] = erlaubt
                    break
        index = spalten.get("perspektivgruppe")
        if index is not None and index < len(felder) and felder[index]:
            gruppen.setdefault(felder[index], []).append(kennung)

    if not quellenarten and not gruppen:
        # Altbestand ohne normierte Tabelle: nur die Quellenart je Zeile lesen.
        for zeile in zeilen:
            kennungen = RE_KENNUNG_IM_TEXT.findall(zeile)
            if len(kennungen) != 1:
                continue
            unten = zeile.casefold()
            for art in QUELLENARTEN:
                if art in unten:
                    quellenarten[kennungen[0]] = art
                    break
    return quellenarten, gruppen


def lade_json(pfad, bezeichnung="Datei"):
    if not pfad:
        return None
    if not os.path.isfile(pfad):
        raise K.KernFehler(f"{bezeichnung} nicht gefunden: {pfad}")
    with open(pfad, encoding="utf-8") as datei:
        return json.load(datei)


def stellen_laden(pfad, artefakt):
    if pfad.lower().endswith((".pptx", ".potx")):
        return stellen_aus_pptx(pfad, artefakt)
    return stellen_aus_markdown(pfad, artefakt)


# ---------------------------------------------------------------------------
# Hauptprogramm
# ---------------------------------------------------------------------------

def nicht_geprueft(args, stellen, baseline, aktiv) -> list[str]:
    """Listet die Gates, die dieser Lauf mangels Eingabe nicht geprüft hat.

    Ein grün protokolliertes Gate, das nichts geprüft hat, ist gefährlicher als
    gar keines. In Modus A trifft das die artefaktgebundenen Gates: Ohne
    Ergebnisdokument haben sie kein Prüfobjekt und liefern zwangsläufig null
    Befunde. Der Bericht muss das benennen, sonst liest sich ein grüner Lauf in
    Modus A wie eine Vollprüfung.
    """
    offen: list[str] = []
    laeuft = lambda name: name in aktiv or args.gate == "alle"  # noqa: E731

    if not stellen:
        if laeuft("zitat"):
            offen.append("Zitat-Gate: nicht anwendbar ohne --deck oder --report. "
                         "Kein Zitat des Ergebnisdokuments wurde gegen den "
                         "*-Pool gehalten.")
        if laeuft("kongruenz"):
            offen.append("Kongruenz-Lint: nicht anwendbar ohne --deck oder "
                         "--report. Zahl-Text-Kongruenz wird nur im Fließtext "
                         "der Endartefakte geprüft.")
        if laeuft("zaehl"):
            offen.append("Zähl-Gate, Dokumentteil: nicht anwendbar ohne --deck "
                         "oder --report. Geprüft wurde allein die Auszählung "
                         "der Arbeitsmappe, keine Zählzeile eines Dokuments.")
        if laeuft("konsistenz"):
            offen.append("Konsistenz-Gate, Dokumentteil: nicht anwendbar ohne "
                         "--deck oder --report. Kategoriennamen, "
                         "Überblicksfolien, abgebrochene Sätze und Typografie "
                         "leben im Ergebnisdokument.")
    if laeuft("wortlaut") and not baseline:
        offen.append("Wortlaut-Sperre: nicht anwendbar ohne --baseline. "
                     "Prosaänderungen gegenüber dem freigegebenen Stand wurden "
                     "nicht verglichen.")
    if laeuft("anonym") and not args.denylist:
        offen.append("Denylist-Abgleich: nicht geprüft ohne --denylist. Der "
                     "blockierende Teil des Anonymisierungs-Gates entfällt.")
    if laeuft("konsistenz") and not args.key:
        offen.append("Registry-Gegenprüfung und Quellenart-Regel: nicht geprüft "
                     "ohne --key.")
    if laeuft("konsistenz") and not args.glossar:
        offen.append("Umbenennungs-Lint: nicht geprüft ohne --glossar.")
    if laeuft("zitat") and stellen and not args.zitatregistry:
        offen.append("Zitat-Registry: nicht geprüft ohne --zitatregistry. Der "
                     "Fuzzy-Match wurde gegen keine deklarierte Quellzeile "
                     "gehalten.")
    return offen


def main(argv=None) -> int:
    parser = argparse.ArgumentParser(
        prog="validate_deliverables.py",
        description="Prüft die Endartefakte gegen die Arbeitsmappe.",
        epilog=HILFE,
        formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--workbook", required=True,
                        help="Arbeitsmappe .xlsx; sie ist die Referenz für alle "
                             "Vergleiche (Pflicht)")
    parser.add_argument("--deck", default=None, help="Ergebnisdokument .pptx")
    parser.add_argument("--report", default=None,
                        help="weiteres Textartefakt, etwa die Prozessnotiz .md")
    parser.add_argument("--denylist", default=None,
                        help="Klarnamen, Orte, Einheiten; ein Treffer ist blockierend")
    parser.add_argument("--key", default=None,
                        help="Pseudonymschlüssel .md oder die JSON-Codiertabelle; "
                             "liefert Quellenart und Gruppenzuordnung")
    parser.add_argument("--glossar", default=None,
                        help="JSON mit verbindlichen Begriffen und Kategoriennamen")
    parser.add_argument("--zitatregistry", default=None,
                        help="JSON mit der beim Bauen notierten Quellzeile je Zitat; "
                             "prüft den Fuzzy-Match gegen die deklarierte Zeile")
    parser.add_argument("--baseline", default=None,
                        help="freigegebener Stand; aktiviert die Wortlaut-Sperre")
    parser.add_argument("--gate", default="alle", choices=GATES,
                        help="einzelnes Gate statt aller (Default: alle)")
    parser.add_argument("--json", dest="json_pfad", default=None,
                        help="Dateipfad für das maschinenlesbare Prüfprotokoll")
    args = parser.parse_args(argv)

    # Artefaktgebundene Gates ohne Endartefakt liefen bisher still mit null
    # Befunden durch. Ein grün protokolliertes Gate, das nichts geprüft hat, ist
    # gefährlicher als gar keines.
    if args.gate in ("zitat", "kongruenz", "wortlaut") and not (args.deck or args.report):
        sys.stderr.write(
            f"FEHLER: Das Gate {args.gate!r} prüft ein Endartefakt gegen die "
            "Arbeitsmappe und braucht deshalb --deck oder --report.\n")
        return 2
    if args.gate == "wortlaut" and not args.baseline:
        sys.stderr.write(
            "FEHLER: Das Gate 'wortlaut' vergleicht gegen den freigegebenen "
            "Stand und braucht deshalb --baseline.\n")
        return 2

    try:
        kontext = baue_kontext(args.workbook)
        denylist = lade_denylist(args.denylist)
        quellenarten, schluesselgruppen = lade_schluessel(args.key)
        glossar = lade_json(args.glossar, "Glossar")
        zitatregistry = lade_json(args.zitatregistry, "Zitat-Registry")
        stellen: list[Stelle] = []
        if args.deck:
            stellen += stellen_laden(args.deck, "Deck")
        if args.report:
            stellen += stellen_laden(args.report, "Bericht")
        baseline = stellen_laden(args.baseline, "Baseline") if args.baseline else None
    except (K.KernFehler, json.JSONDecodeError) as fehler:
        sys.stderr.write(f"FEHLER: {fehler}\n")
        return 2

    modus = "A (nur Arbeitsmappe)" if not stellen else "B (Endartefakte gegen Arbeitsmappe)"
    aktiv = GATES if args.gate == "alle" else (args.gate,)
    befunde: list[K.Befund] = []

    if args.gate in ("alle", "zaehl", "konsistenz"):
        befunde += gate_arbeitsmappe(kontext)
    if stellen:
        if "zaehl" in aktiv or args.gate == "alle":
            befunde += gate_zaehl(stellen, kontext)
        if "zitat" in aktiv or args.gate == "alle":
            befunde += gate_zitat(stellen, kontext, quellenarten, zitatregistry)
        if "konsistenz" in aktiv or args.gate == "alle":
            befunde += gate_konsistenz(stellen, kontext, glossar)
        if "kongruenz" in aktiv or args.gate == "alle":
            befunde += gate_kongruenz(stellen, kontext)
    if "anonym" in aktiv or args.gate == "alle":
        if denylist:
            befunde += gate_anonym(stellen, kontext, denylist)
        else:
            befunde += gate_anonym(stellen, kontext, [])
            befunde.append(K.Befund(
                "anonym", "warnung", "Aufruf",
                "Denylist-Gate nicht geprüft.",
                "Ohne --denylist läuft der blockierende Abgleich gegen Klarnamen, "
                "Orte und Einheitenbezeichnungen nicht. Lege die Denylist aus dem "
                "Pseudonymschlüssel an und übergib sie."))
    if baseline and ("wortlaut" in aktiv or args.gate == "alle"):
        befunde += gate_wortlaut(stellen, baseline)

    if "konsistenz" in aktiv or args.gate == "alle":
        befunde += gate_registry(kontext, args.key, schluesselgruppen)

    fehlerzahl = sum(1 for b in befunde if b.blockierend)
    warnzahl = len(befunde) - fehlerzahl

    print("# Prüfprotokoll\n")
    print(f"- Modus: {modus}")
    print(f"- Arbeitsmappe: `{os.path.basename(args.workbook)}`")
    if args.deck:
        print(f"- Ergebnisdokument: `{os.path.basename(args.deck)}`")
    if args.report:
        print(f"- Bericht: `{os.path.basename(args.report)}`")
    print(f"- Gate: {args.gate}")
    print(f"- Datenbasis: {kontext['einheiten']} Codiereinheiten, "
          f"N = {kontext['n']}, {len(kontext['kategorien'])} Unterkategorien")
    print(f"- Zitatpool: {sum(1 for e in kontext['pool'] if e['stern'])} "
          f"von {kontext['einheiten']} Einheiten mit *\n")

    print(K.bericht_markdown("Befunde", befunde))

    protokoll = kontext.get("zitatprotokoll") or []
    if protokoll:
        print("## Zitat-Gate im Einzelnen\n")
        print("| Ort | Status | Abweichung | Wert | Quelle | im *-Pool | Marker "
              "| Bewertung |")
        print("|---|---|---|---:|---|---|---|---|")
        for eintrag in protokoll:
            print(f"| {eintrag['ort']} | {eintrag['status']} | "
                  f"{eintrag['abweichung'] or 'keine'} | "
                  f"{eintrag['aehnlichkeit']:.2f} | {eintrag['quelle'] or 'keine'} | "
                  f"{'ja' if eintrag['im_sternpool'] else 'nein'} | "
                  f"{eintrag['marker'] or 'ohne'} | {eintrag['bewertung'] or 'ohne'} |")
        print()

    print("## Kennzahlen\n")
    print("| Kennzahl | Wert |")
    print("|---|---:|")
    print(f"| geprüfte Textstellen | {len(stellen)} |")
    print(f"| geprüfte Zitate | {len(protokoll)} |")
    print(f"| blockierende Verstöße | {fehlerzahl} |")
    print(f"| Warnungen | {warnzahl} |")
    print()

    offen = nicht_geprueft(args, stellen, baseline, aktiv)
    print("## Nicht geprüft\n")
    if offen:
        for eintrag in offen:
            print(f"- {eintrag}")
        print("\nEin grünes Ergebnis deckt nur, was oben tatsächlich lief.")
    else:
        print("- nichts; alle aktiven Gates hatten ihr Prüfobjekt.")
    print()
    print(f"Ergebnis: {K.bilanz(fehlerzahl, warnzahl)}.")

    if args.json_pfad:
        ausgabe = {
            "workbook": args.workbook, "deck": args.deck, "report": args.report,
            "modus": modus, "gate": args.gate,
            "einheiten": kontext["einheiten"], "n_personen": kontext["n"],
            "befunde": [b.als_dict() for b in befunde],
            "zitate": protokoll,
            "nicht_geprueft": offen,
            "blockierend": fehlerzahl, "warnungen": warnzahl,
        }
        try:
            with open(args.json_pfad, "w", encoding="utf-8") as datei:
                json.dump(ausgabe, datei, ensure_ascii=False, indent=2)
        except OSError as fehler:
            sys.stderr.write(f"FEHLER: Prüfprotokoll nicht schreibbar: {fehler}\n")
            return 2

    return 1 if fehlerzahl else 0


if __name__ == "__main__":
    raise SystemExit(main())
